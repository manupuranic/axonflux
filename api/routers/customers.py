import csv
import io
from datetime import date

import openpyxl
import openpyxl.utils
from openpyxl.styles import Alignment, Font, PatternFill

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import text
from sqlalchemy.engine import Connection

from api.dependencies import get_conn, get_current_user
from api.lib.filters import FieldSpec, build_where, parse_conditions
from api.schemas.auth import CurrentUser
from api.schemas.customers import (
    ChurnTier,
    CustomerBill,
    CustomerListItem,
    CustomerSummary,
    LapsedCustomer,
    LapsedSummary,
)

router = APIRouter(prefix="/api/customers", tags=["customers"])


@router.get("", response_model=dict)
def list_customers(
    search: str = Query(default=None, max_length=100),
    is_repeat: bool = Query(default=None),
    is_member: bool = Query(default=None),
    include_walkin: bool = Query(default=False),
    sort_by: str = Query(default="total_revenue", pattern="^(total_revenue|total_bills|last_seen_date|avg_bill_value)$"),
    limit: int = Query(default=50, le=200),
    offset: int = Query(default=0, ge=0),
    conn: Connection = Depends(get_conn),
    _: CurrentUser = Depends(get_current_user),
):
    """Paginated customer list, sorted by revenue by default. Walk-ins excluded unless requested."""
    filters = []
    params: dict = {}

    if not include_walkin:
        filters.append("d.is_walk_in = FALSE")

    if search:
        filters.append("(d.display_name ILIKE :search OR d.mobile_clean ILIKE :search)")
        params["search"] = f"%{search}%"

    if is_repeat is not None:
        filters.append("m.is_repeat = :is_repeat")
        params["is_repeat"] = is_repeat

    if is_member is not None:
        filters.append("d.is_member = :is_member")
        params["is_member"] = is_member

    where_clause = ("WHERE " + " AND ".join(filters)) if filters else ""

    sort_map = {
        "total_revenue": "m.total_revenue DESC NULLS LAST",
        "total_bills": "m.total_bills DESC NULLS LAST",
        "last_seen_date": "d.last_seen_date DESC NULLS LAST",
        "avg_bill_value": "m.avg_bill_value DESC NULLS LAST",
    }
    order_clause = sort_map[sort_by]

    base_sql = f"""
        FROM derived.customer_dimension d
        LEFT JOIN derived.customer_metrics m ON d.mobile_clean = m.mobile_clean
        {where_clause}
    """

    total = conn.execute(text(f"SELECT COUNT(*) {base_sql}"), params).scalar()

    rows = conn.execute(
        text(f"""
            SELECT
                d.mobile_clean,
                d.display_name,
                d.is_walk_in,
                d.is_member,
                d.first_seen_date,
                d.last_seen_date,
                d.total_bills,
                m.total_revenue,
                m.avg_bill_value,
                m.total_discount_received,
                m.days_since_last_visit,
                m.avg_days_between_visits,
                m.is_repeat,
                m.preferred_payment
            {base_sql}
            ORDER BY {order_clause}
            LIMIT :limit OFFSET :offset
        """),
        {**params, "limit": limit, "offset": offset},
    ).mappings().all()

    return {
        "total": total,
        "limit": limit,
        "offset": offset,
        "items": [CustomerListItem(**dict(r)) for r in rows],
    }


@router.get("/summary", response_model=CustomerSummary)
def get_customer_summary(
    conn: Connection = Depends(get_conn),
    _: CurrentUser = Depends(get_current_user),
):
    """Aggregate KPIs across all identified (non-walk-in) customers."""
    row = conn.execute(text("""
        WITH identified AS (
            SELECT d.mobile_clean, d.is_member, d.first_seen_date, m.avg_bill_value, m.is_repeat
            FROM derived.customer_dimension d
            LEFT JOIN derived.customer_metrics m ON d.mobile_clean = m.mobile_clean
            WHERE d.is_walk_in = FALSE
        ),
        total_rev AS (
            SELECT SUM(total_revenue) AS total
            FROM derived.customer_metrics
        ),
        walkin_rev AS (
            SELECT COALESCE(total_revenue, 0) AS total
            FROM derived.customer_metrics
            WHERE mobile_clean = 'WALK-IN'
        )
        SELECT
            COUNT(*)                                                          AS total_unique_customers,
            COUNT(*) FILTER (WHERE is_repeat = TRUE)                         AS repeat_customer_count,
            ROUND(
                COUNT(*) FILTER (WHERE is_repeat = TRUE) * 100.0
                / NULLIF(COUNT(*), 0), 1
            )                                                                 AS repeat_customer_percent,
            AVG(avg_bill_value)                                               AS avg_bill_value,
            COUNT(*) FILTER (WHERE is_member = TRUE)                         AS members_count,
            COUNT(*) FILTER (WHERE first_seen_date >= CURRENT_DATE - 30)     AS new_customers_last_30d,
            ROUND(
                (SELECT total FROM walkin_rev) * 100.0
                / NULLIF((SELECT total FROM total_rev), 0), 1
            )                                                                 AS walk_in_revenue_percent
        FROM identified
    """)).mappings().one()

    return CustomerSummary(
        total_unique_customers=int(row["total_unique_customers"] or 0),
        repeat_customer_count=int(row["repeat_customer_count"] or 0),
        repeat_customer_percent=float(row["repeat_customer_percent"] or 0),
        avg_bill_value=float(row["avg_bill_value"]) if row["avg_bill_value"] else None,
        members_count=int(row["members_count"] or 0),
        new_customers_last_30d=int(row["new_customers_last_30d"] or 0),
        walk_in_revenue_percent=float(row["walk_in_revenue_percent"]) if row["walk_in_revenue_percent"] else None,
    )



_LAPSED_BASE_SQL = """
    FROM derived.customer_dimension d
    JOIN derived.customer_metrics m ON d.mobile_clean = m.mobile_clean
    WHERE d.is_walk_in = FALSE
      AND m.is_repeat = TRUE
"""

_CHURN_TIER_EXPR = """
    CASE
        WHEN m.days_since_last_visit >= 90 THEN 'lost'
        WHEN m.days_since_last_visit >= 60 THEN 'lapsed'
        WHEN m.days_since_last_visit >= 30 THEN 'at-risk'
        ELSE 'active'
    END
"""


def _tier_filter(tier: str | None) -> str:
    if tier == "active":
        return "AND m.days_since_last_visit < 30"
    if tier == "at-risk":
        return "AND m.days_since_last_visit >= 30 AND m.days_since_last_visit < 60"
    if tier == "lapsed":
        return "AND m.days_since_last_visit >= 60 AND m.days_since_last_visit < 90"
    if tier == "lost":
        return "AND m.days_since_last_visit >= 90"
    return ""


# Public field spec for the Customer Activity endpoint. Keys are wire-format
# names the client sends in `cond=key:op:value`. Columns stay server-internal,
# never leak to the URL. Adding a new filterable column = one line here.
LAPSED_FIELDS: dict[str, FieldSpec] = {
    "visits":       FieldSpec(column="m.total_bills",            type="number"),
    "days_silent":  FieldSpec(column="m.days_since_last_visit",  type="number"),
    "spend":        FieldSpec(column="m.total_revenue",          type="number"),
    "avg_bill":     FieldSpec(column="m.avg_bill_value",         type="number"),
    "name":         FieldSpec(column="d.display_name",           type="string"),
    "mobile":       FieldSpec(column="d.mobile_clean",           type="string"),
    "is_member":    FieldSpec(column="d.is_member",              type="bool"),
    "payment":      FieldSpec(
        column="m.preferred_payment",
        type="enum",
        allowed_values=("cash", "card", "upi", "credit"),
    ),
}


def _parse_lapsed_conditions(cond: list[str]) -> tuple[str, dict]:
    try:
        filters = parse_conditions(cond, LAPSED_FIELDS)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"Invalid filter: {exc}") from exc
    return build_where(filters)


@router.get("/lapsed", response_model=dict)
def list_lapsed_customers(
    tier: ChurnTier | None = Query(default=None),
    cond: list[str] = Query(default_factory=list, description="Filter as key:op:value (repeatable)"),
    limit: int = Query(default=50, le=500),
    offset: int = Query(default=0, ge=0),
    conn: Connection = Depends(get_conn),
    _: CurrentUser = Depends(get_current_user),
):
    tier_sql = _tier_filter(tier)
    filter_sql, filter_params = _parse_lapsed_conditions(cond)
    base = f"{_LAPSED_BASE_SQL} {tier_sql} {filter_sql}"

    # Summary KPIs intentionally ignore the user's filter knobs — they always
    # reflect the full tier breakdown so the user sees the unfiltered universe
    # before slicing into it. (Matches the existing UX: tier chips show full counts.)
    summary_row = conn.execute(text(f"""
        SELECT
            COUNT(*) FILTER (WHERE m.days_since_last_visit < 30)                                    AS active_count,
            COUNT(*) FILTER (WHERE m.days_since_last_visit >= 30 AND m.days_since_last_visit < 60) AS at_risk_count,
            COUNT(*) FILTER (WHERE m.days_since_last_visit >= 60 AND m.days_since_last_visit < 90) AS lapsed_count,
            COUNT(*) FILTER (WHERE m.days_since_last_visit >= 90)                                   AS lost_count,
            COUNT(*)                                                                                 AS total_count
        {_LAPSED_BASE_SQL}
    """)).mappings().one()

    total = conn.execute(text(f"SELECT COUNT(*) {base}"), filter_params).scalar()

    rows = conn.execute(text(f"""
        SELECT
            d.mobile_clean,
            d.display_name,
            d.is_member,
            m.total_bills,
            m.total_revenue,
            m.avg_bill_value,
            m.last_purchase_date,
            m.days_since_last_visit,
            m.avg_days_between_visits,
            m.preferred_payment,
            {_CHURN_TIER_EXPR} AS churn_tier
        {base}
        ORDER BY m.days_since_last_visit DESC
        LIMIT :limit OFFSET :offset
    """), {**filter_params, "limit": limit, "offset": offset}).mappings().all()

    return {
        "total": total,
        "limit": limit,
        "offset": offset,
        "summary": LapsedSummary(**dict(summary_row)),
        "items": [LapsedCustomer(**dict(r)) for r in rows],
    }


_EXPORT_HEADERS = [
    "Name", "Mobile", "Status", "Last Visit", "Days Silent",
    "Total Spend (₹)", "Visits", "Avg Bill (₹)", "Member", "Pays With",
]


@router.get("/lapsed/export")
def export_lapsed_customers(
    tier: ChurnTier | None = Query(default=None),
    cond: list[str] = Query(default_factory=list),
    export_format: str = Query(default="csv", pattern="^(csv|xlsx)$"),
    conn: Connection = Depends(get_conn),
    _: CurrentUser = Depends(get_current_user),
):
    tier_sql = _tier_filter(tier)
    filter_sql, filter_params = _parse_lapsed_conditions(cond)
    base = f"{_LAPSED_BASE_SQL} {tier_sql} {filter_sql}"

    rows = conn.execute(text(f"""
        SELECT
            d.display_name,
            d.mobile_clean,
            {_CHURN_TIER_EXPR} AS churn_tier,
            m.last_purchase_date,
            m.days_since_last_visit,
            m.total_revenue,
            m.total_bills,
            m.avg_bill_value,
            d.is_member,
            m.preferred_payment
        {base}
        ORDER BY m.days_since_last_visit DESC
        LIMIT 10000
    """), filter_params).mappings().all()

    today = date.today().isoformat()
    tier_str = f"_{tier}" if tier else ""
    filename_base = f"lapsed_customers{tier_str}_{today}"

    def _row_values(r):
        return [
            r["display_name"] or "",
            r["mobile_clean"],
            r["churn_tier"],
            str(r["last_purchase_date"]) if r["last_purchase_date"] is not None else "",
            r["days_since_last_visit"] if r["days_since_last_visit"] is not None else "",
            round(float(r["total_revenue"]), 2) if r["total_revenue"] is not None else "",
            r["total_bills"] if r["total_bills"] is not None else "",
            round(float(r["avg_bill_value"]), 2) if r["avg_bill_value"] is not None else "",
            "Yes" if r["is_member"] else "No",
            (r["preferred_payment"] or "").upper(),
        ]

    if export_format == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(_EXPORT_HEADERS)
        for r in rows:
            writer.writerow(_row_values(r))
        buf.seek(0)
        return StreamingResponse(
            iter([buf.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.csv"'},
        )

    # xlsx
    tier_colors = {"at-risk": "FFF59D", "lapsed": "FFCC80", "lost": "EF9A9A"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lapsed Customers"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    for col_idx, header in enumerate(_EXPORT_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, r in enumerate(rows, 2):
        values = _row_values(r)
        for col_idx, val in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
        tier_val = r["churn_tier"]
        if tier_val in tier_colors:
            fill = PatternFill("solid", fgColor=tier_colors[tier_val])
            ws.cell(row=row_idx, column=3).fill = fill

    col_widths = [24, 16, 12, 14, 13, 18, 8, 14, 8, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    buf_bytes = io.BytesIO()
    wb.save(buf_bytes)
    buf_bytes.seek(0)
    return StreamingResponse(
        iter([buf_bytes.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename_base}.xlsx"'},
    )


@router.get("/{mobile}/history", response_model=list[CustomerBill])
def get_customer_history(
    mobile: str,
    limit: int = Query(default=50, le=200),
    conn: Connection = Depends(get_conn),
    _: CurrentUser = Depends(get_current_user),
):
    """All bills for a customer identified by their normalized 10-digit mobile."""
    rows = conn.execute(
        text("""
            SELECT
                bill_no,
                TO_TIMESTAMP(
                    CASE WHEN SUBSTRING(bill_datetime_raw, 11, 1) = ' '
                         THEN bill_datetime_raw
                         ELSE SUBSTRING(bill_datetime_raw, 1, 10) || ' ' || SUBSTRING(bill_datetime_raw, 11)
                    END,
                    CASE WHEN bill_datetime_raw ~* '(AM|PM)\s*$'
                         THEN 'DD-MM-YYYY HH12:MI AM'
                         ELSE 'DD-MM-YYYY HH24:MI'
                    END
                )::DATE AS bill_date,
                net_total,
                total_discount,
                COALESCE(actual_cash, cash_amount, 0)                                    AS cash_total,
                COALESCE(card_amount, 0)                                                 AS card_total,
                COALESCE(google_pay_amount,0)+COALESCE(phonepe_amount,0)
                    +COALESCE(paytm_amount,0)                                            AS upi_total,
                COALESCE(credit_amount, 0)                                               AS credit_total
            FROM raw.raw_sales_billwise
            WHERE REGEXP_REPLACE(COALESCE(customer_mobile_raw,''), '[^0-9]','','g')
                      ~ ('^(91|0)?'||:mobile||'$')
               OR REGEXP_REPLACE(COALESCE(customer_mobile_raw,''), '[^0-9]','','g') = :mobile
            ORDER BY bill_datetime_raw DESC
            LIMIT :limit
        """),
        {"mobile": mobile, "limit": limit},
    ).mappings().all()

    return [CustomerBill(**dict(r)) for r in rows]
