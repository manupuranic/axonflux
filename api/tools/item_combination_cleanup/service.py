from __future__ import annotations

import os
import uuid
import hashlib
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from sqlalchemy import case, func, or_
from sqlalchemy.engine import Connection
from sqlalchemy.orm import Session

from api.schemas.auth import CurrentUser
from api.tools.item_combination_cleanup.classify import build_loose_anchors, classify
from api.tools.item_combination_cleanup.models import (
    ItemCombinationCleanupRow,
    ItemCombinationCleanupRun,
    ItemCombinationCleanupSemanticAssessment,
    ItemCombinationCleanupFieldDecision,
    ItemCombinationCleanupNameSuggestion,
)
from api.tools.item_combination_cleanup.naming_standard import suggest_name_standard
from api.tools.item_combination_cleanup.normalize import propose_name_and_size
from api.tools.item_combination_cleanup.purchases import (
    barcode_key,
    load_purchase_index,
    mrp_key,
)
from api.tools.item_combination_cleanup.workbook import (
    WorkbookError,
    WorkbookValidationError,
    export_identity_copy,
    inspect_workbook,
    iter_source_rows,
    _item_id_key,
    export_approved_proposals,
)

_MAX_UPLOAD_BYTES = 25 * 1024 * 1024
_ALLOWED_SUFFIXES = {".xlsx"}
_PROTECTED_IDENTITY_KEYS = (
    "item_id",
    "barcode",
    "mrp",
    "rate",
    "purchase_price",
    "expiry",
    "stock",
)


class CleanupServiceError(Exception):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def runs_root() -> Path:
    override = os.getenv("ITEM_COMBINATION_CLEANUP_DIR")
    if override:
        return Path(override)
    return Path(__file__).resolve().parents[3] / "data" / "uploads" / "item_combination_cleanup"


def run_dir(run_id: uuid.UUID) -> Path:
    return runs_root() / str(run_id)


def _validate_upload_name(filename: str | None) -> str:
    name = Path(filename or "").name
    if not name or name in {".", ".."}:
        raise CleanupServiceError("A file name is required.", 400)
    if Path(name).suffix.lower() not in _ALLOWED_SUFFIXES:
        raise CleanupServiceError("Only .xlsx files are accepted.", 400)
    return name


def _json_safe(value: object) -> object:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    return str(value)


def _as_text(value: object | None) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if text.strip() != "" else None


def create_run(
    db: Session,
    *,
    filename: str | None,
    content: bytes,
    content_type: str | None,
    user: CurrentUser,
) -> ItemCombinationCleanupRun:
    safe_name = _validate_upload_name(filename)
    _ = content_type
    if len(content) > _MAX_UPLOAD_BYTES:
        raise CleanupServiceError("File exceeds 25 MB limit.", 413)
    if not content.startswith(b"PK"):
        raise CleanupServiceError("File is not a valid XLSX workbook.", 400)

    run_id = uuid.uuid4()
    dest_dir = run_dir(run_id)
    dest_dir.mkdir(parents=True, exist_ok=True)
    source_path = dest_dir / "source.xlsx"
    source_path.write_bytes(content)

    try:
        meta = inspect_workbook(source_path)
    except WorkbookError as exc:
        source_path.unlink(missing_ok=True)
        try:
            dest_dir.rmdir()
        except OSError:
            pass
        raise CleanupServiceError(str(exc), 422) from exc

    run = ItemCombinationCleanupRun(
        id=run_id,
        created_by=uuid.UUID(user.id) if user.id else None,
        status="uploaded",
        source_file_name=safe_name,
        original_path=str(source_path),
        sheet_name=meta.sheet_name,
        header_row=meta.header_row,
        row_count=meta.data_row_count,
        column_count=meta.column_count,
        headers=list(meta.headers),
        validation_status=None,
        error_message=None,
    )
    db.add(run)
    db.flush()
    return run


def list_runs(db: Session, limit: int = 50) -> list[ItemCombinationCleanupRun]:
    return (
        db.query(ItemCombinationCleanupRun)
        .order_by(ItemCombinationCleanupRun.created_at.desc())
        .limit(limit)
        .all()
    )


def get_run(db: Session, run_id: uuid.UUID) -> ItemCombinationCleanupRun | None:
    return db.query(ItemCombinationCleanupRun).filter(ItemCombinationCleanupRun.id == run_id).first()


def export_run(db: Session, run: ItemCombinationCleanupRun) -> ItemCombinationCleanupRun:
    source = Path(run.original_path)
    if not source.is_file():
        raise CleanupServiceError("Original workbook is missing on disk.", 404)

    dest = run_dir(run.id) / "output.xlsx"
    try:
        export_identity_copy(source, dest)
    except WorkbookValidationError as exc:
        run.validation_status = "failed"
        run.error_message = str(exc)
        run.output_path = None
        db.flush()
        raise CleanupServiceError(str(exc), 422) from exc

    run.output_path = str(dest)
    run.status = "exported"
    run.validation_status = "passed"
    run.error_message = None
    db.flush()
    return run


def process_run(
    db: Session,
    conn: Connection,
    run: ItemCombinationCleanupRun,
) -> ItemCombinationCleanupRun:
    """Recompute proposals from the original workbook + raw purchases.

    Idempotent: previous proposal rows for the run are replaced.
    Does not mutate source.xlsx or write output.xlsx.
    """
    source = Path(run.original_path)
    if not source.is_file():
        raise CleanupServiceError("Original workbook is missing on disk.", 404)
    if db.query(ItemCombinationCleanupRow).filter(
        ItemCombinationCleanupRow.run_id == run.id,
        ItemCombinationCleanupRow.approval_status != "PENDING",
    ).count():
        raise CleanupServiceError("Reviewed runs cannot be reprocessed; create a new run.", 409)

    _, source_rows = iter_source_rows(source)
    purchases = load_purchase_index(conn)

    prepared: list[dict] = []
    bm_item_ids: dict[tuple[str, Decimal], set[str]] = defaultdict(set)
    barcode_item_identity: dict[str, list[tuple[str | None, str | None]]] = defaultdict(list)
    for row in source_rows:
        item_key = _item_id_key(row.item_id)
        if item_key is None:
            continue
        ns = propose_name_and_size(row.item_name, row.size)
        bc = barcode_key(row.barcode)
        mrp = mrp_key(row.mrp)
        hit = purchases.lookup(bc, mrp)
        if bc and mrp is not None:
            bm_item_ids[(bc, mrp)].add(item_key)
        if bc:
            barcode_item_identity[bc].append((_as_text(row.item_name), _as_text(row.size)))
        prepared.append(
            {
                "row": row,
                "item_key": item_key,
                "ns": ns,
                "bc": bc or "",
                "mrp": mrp,
                "hit": hit,
            }
        )

    anchors = build_loose_anchors(
        [
            (p["item_key"], p["bc"], p["ns"].proposed_name, p["hit"] is not None)
            for p in prepared
        ]
    )

    db.query(ItemCombinationCleanupRow).filter(
        ItemCombinationCleanupRow.run_id == run.id
    ).delete(synchronize_session=False)
    db.flush()

    type_counts: Counter[str] = Counter()
    confidence_counts: Counter[str] = Counter()
    review_counts: Counter[str] = Counter()

    for p in prepared:
        row = p["row"]
        ns = p["ns"]
        hit = p["hit"]
        bm_count = 0
        if p["bc"] and p["mrp"] is not None:
            bm_count = len(bm_item_ids[(p["bc"], p["mrp"])])
        source_identity = barcode_item_identity[p["bc"]]
        barcode_identity_safe = purchases.identity_is_safe(
            p["bc"],
            item_names=[name for name, _ in source_identity],
            item_sizes=[size for _, size in source_identity],
        )
        clf = classify(
            normalized_name=ns.proposed_name,
            original_brand=_as_text(row.brand),
            purchase=hit,
            barcode_mrp_item_id_count=bm_count,
            loose_anchors=anchors,
            barcode_identity_safe=barcode_identity_safe,
        )
        review_status = clf.review_status
        if clf.product_type == "UNKNOWN" and ns.evidence:
            review_status = "AUTO_APPROVED"

        identity = {
            "item_id": _json_safe(row.item_id),
            "barcode": _json_safe(row.barcode),
            "mrp": _json_safe(row.mrp),
            "rate": _json_safe(row.rate),
            "purchase_price": _json_safe(row.purchase_price),
            "expiry": _json_safe(row.expiry),
            "stock": _json_safe(row.stock),
        }
        db.add(
            ItemCombinationCleanupRow(
                run_id=run.id,
                row_index=row.row_index,
                item_id_key=p["item_key"],
                source_identity=identity,
                original_item_name=_as_text(row.item_name),
                original_brand=_as_text(row.brand),
                original_size=_as_text(row.size),
                original_hsn=_as_text(row.hsn),
                proposed_item_name=ns.proposed_name or _as_text(row.item_name),
                proposed_size=ns.proposed_size,
                proposed_brand=clf.proposed_brand,
                product_type=clf.product_type,
                classification_confidence=clf.confidence,
                supplier_name=hit.supplier_name if hit else None,
                supplier_purchase_date=hit.purchase_date if hit else None,
                supplier_purchase_id=hit.purchase_id if hit else None,
                supplier_invoice_no=hit.invoice_no if hit else None,
                supplier_source_file=hit.source_file_name if hit else None,
                supplier_match_method=hit.match_method if hit else None,
                classification_evidence=_json_safe(
                    {**clf.evidence, "rules": clf.rules}
                ),
                name_evidence=_json_safe(ns.evidence),
                review_status=review_status,
                approval_status="PENDING",
            )
        )
        type_counts[clf.product_type] += 1
        if clf.confidence:
            confidence_counts[clf.confidence] += 1
        review_counts[review_status] += 1

    run.summary_json = {
        "row_count": len(prepared),
        "product_type": dict(type_counts),
        "confidence": dict(confidence_counts),
        "review_status": dict(review_counts),
        "protected_identity_keys": list(_PROTECTED_IDENTITY_KEYS),
    }
    run.status = "processed"
    run.error_message = None
    db.flush()
    return run


def _text_changed(left: str | None, right: str | None) -> bool:
    return (left or "").strip() != (right or "").strip()


def proposed_changes(row) -> dict[str, dict[str, str | None]]:
    """Response-only change set. None means no proposal; empty string means clear."""
    fields = (
        ("name", row.original_item_name, row.proposed_item_name),
        ("brand", row.original_brand, row.proposed_brand),
        ("size", row.original_size, row.proposed_size),
    )
    return {
        key: {"original": original, "proposed": proposed}
        for key, original, proposed in fields
        if proposed is not None and _text_changed(original, proposed)
    }


def resolved_reviewed_value(original: str | None, proposed: str | None, decision: str | None, edited: str | None) -> str | None:
    """An edit is provisional until an explicit ACCEPTED decision confirms it."""
    if decision == "ACCEPTED":
        return edited if edited is not None else proposed
    return original


def effective_name_for_phase5(db: Session, row: ItemCombinationCleanupRow) -> str | None:
    """Mirror export precedence without making Phase 5 suggestions export-authoritative."""
    decision = db.query(ItemCombinationCleanupFieldDecision).filter_by(run_id=row.run_id, item_id_key=row.item_id_key, field_name="name").first()
    if decision:
        return resolved_reviewed_value(
            row.original_item_name,
            getattr(decision, "deterministic_proposed_value", None) or row.proposed_item_name,
            decision.decision,
            decision.edited_value,
        )
    # A row-level approval (for example a Brand or Size decision) is not name
    # authority. Phase 5 may only inherit an explicit accepted name decision.
    return row.original_item_name


def generate_phase5_name_suggestions(db: Session, run: ItemCombinationCleanupRun) -> dict[str, int]:
    """Idempotently persist additive suggestions; never changes existing proposals or decisions."""
    created = unchanged = 0
    existing = {(s.item_id_key, s.suggestion_version): s for s in db.query(ItemCombinationCleanupNameSuggestion).filter_by(run_id=run.id, suggestion_source="NAMING_STANDARD")}
    for row in db.query(ItemCombinationCleanupRow).filter_by(run_id=run.id):
        base = effective_name_for_phase5(db, row)
        suggestion = suggest_name_standard(base)
        if not suggestion.suggested_name:
            continue
        key = (row.item_id_key, "phase5b-v1")
        if key in existing:
            unchanged += 1
            continue
        db.add(ItemCombinationCleanupNameSuggestion(run_id=run.id, item_id_key=row.item_id_key, field_name="name", suggestion_source="NAMING_STANDARD", suggestion_version="phase5b-v1", category=suggestion.category or "ABBREVIATION_EXPANSION", base_value=base or "", suggested_value=suggestion.suggested_name, transformations=suggestion.transformations, evidence=suggestion.evidence))
        created += 1
    db.flush()
    return {"created": created, "existing": unchanged}


_STALE_NAME_SUGGESTION_STATUS = "SUPERSEDED_STALE_BASELINE"
_BULK_SAFE_NAME_TRANSFORMATIONS = {
    "CHOCOLT → CHOCOLATE",
    "SHOULDE → SHOULDERS",
    "LIQUD → LIQUID",
    "CHESE → CHEESE",
    "WIPERC → WIPER",
}


def _name_decision_supersedes_suggestion(decision, suggestion) -> bool:
    if not decision or decision.decision not in {"ACCEPTED", "REJECTED"}:
        return False
    if decision.reviewed_at is None or suggestion.created_at is None:
        return True
    return decision.reviewed_at >= suggestion.created_at


def _name_overlap_classification(suggestions, decision=None) -> str | None:
    if decision and any(_name_decision_supersedes_suggestion(decision, s) for s in suggestions):
        return "ALREADY_RESOLVED_BY_HUMAN"
    deterministic = next((s for s in suggestions if s.suggestion_source == "NAMING_STANDARD"), None)
    semantic = next((s for s in suggestions if s.suggestion_source == "SEMANTIC_V2"), None)
    if not deterministic or not semantic:
        return None
    if deterministic.suggested_value == semantic.suggested_value:
        if deterministic.base_value == semantic.base_value:
            return "EXACT_SAME_RESULT"
        return "DETERMINISTIC_SUPERSEDES_SEMANTIC"

    det_originals = {
        value.split(" → ", 1)[0]
        for value in (deterministic.transformations or [])
        if isinstance(value, str) and " → " in value
    }
    semantic_originals = {
        value.get("original_text")
        for value in (semantic.transformations or [])
        if isinstance(value, dict) and value.get("original_text")
    }
    if det_originals and semantic_originals and det_originals.isdisjoint(semantic_originals):
        return "SEMANTIC_ADDS_ANOTHER_CHANGE"
    return "CONFLICTING_RESULTS"


def _name_suggestion_out(suggestion, *, safe_bulk_group: str | None = None) -> dict:
    evidence = suggestion.evidence or {}
    payload = evidence.get("payload") or {}
    candidate = payload.get("candidate") or {}
    context = payload.get("context") or {}
    return {
        "id": str(suggestion.id),
        "source": suggestion.suggestion_source,
        "version": suggestion.suggestion_version,
        "category": suggestion.category,
        "base_value": suggestion.base_value,
        "suggested_value": suggestion.suggested_value,
        "transformations": suggestion.transformations or [],
        "reason": evidence.get("reason"),
        "confidence": evidence.get("confidence"),
        "detector_reason": candidate.get("detector_reason"),
        "corroborating_evidence": candidate.get("corroborating_evidence"),
        "purchase_item_names": context.get("purchase_item_names") or [],
        "catalog_siblings": context.get("catalog_siblings") or [],
        "safe_bulk_group": safe_bulk_group,
        "status": suggestion.status,
    }


def list_actionable_name_review(db: Session, run_id: uuid.UUID) -> dict:
    """Read-only Phase 5 review projection; suggestion rows remain advisory."""
    rows = {
        row.item_id_key: row
        for row in db.query(ItemCombinationCleanupRow).filter_by(run_id=run_id)
    }
    decisions = {
        decision.item_id_key: decision
        for decision in db.query(ItemCombinationCleanupFieldDecision).filter_by(
            run_id=run_id, field_name="name"
        )
    }
    suggestions = db.query(ItemCombinationCleanupNameSuggestion).filter_by(run_id=run_id).all()
    actionable = []
    stale_items = set()
    decided_items = set()
    for suggestion in suggestions:
        row = rows.get(suggestion.item_id_key)
        if row is None:
            stale_items.add(suggestion.item_id_key)
            continue
        if suggestion.status == _STALE_NAME_SUGGESTION_STATUS:
            stale_items.add(suggestion.item_id_key)
            continue
        if suggestion.status != "PENDING":
            decided_items.add(suggestion.item_id_key)
            continue
        decision = decisions.get(suggestion.item_id_key)
        if _name_decision_supersedes_suggestion(decision, suggestion):
            decided_items.add(suggestion.item_id_key)
            continue
        effective = effective_name_for_phase5(db, row) or ""
        if (
            not suggestion.base_value
            or not suggestion.suggested_value
            or suggestion.base_value == suggestion.suggested_value
            or suggestion.suggested_value == effective
            or suggestion.base_value != effective
        ):
            stale_items.add(suggestion.item_id_key)
            continue
        actionable.append(suggestion)

    by_item = defaultdict(list)
    for suggestion in actionable:
        by_item[suggestion.item_id_key].append(suggestion)

    items = []
    deterministic_only = semantic_only = overlap = 0
    for item_id, item_suggestions in by_item.items():
        sources = {suggestion.suggestion_source for suggestion in item_suggestions}
        if sources == {"NAMING_STANDARD"}:
            deterministic_only += 1
        elif sources == {"SEMANTIC_V2"}:
            semantic_only += 1
        else:
            overlap += 1
        safe_group = None
        if sources == {"NAMING_STANDARD"} and len(item_suggestions) == 1:
            transformations = item_suggestions[0].transformations or []
            if len(transformations) == 1 and transformations[0] in _BULK_SAFE_NAME_TRANSFORMATIONS:
                safe_group = transformations[0]
        row = rows[item_id]
        decision = decisions.get(item_id)
        items.append(
            {
                "item_id": item_id,
                "effective_name": effective_name_for_phase5(db, row),
                "overlap_classification": _name_overlap_classification(item_suggestions, decision),
                "current_name_decision": None if decision is None else {
                    "decision": decision.decision,
                    "edited_value": decision.edited_value,
                    "reviewed_at": decision.reviewed_at,
                },
                "suggestions": [
                    _name_suggestion_out(suggestion, safe_bulk_group=safe_group)
                    for suggestion in sorted(
                        item_suggestions,
                        key=lambda value: 0 if value.suggestion_source == "NAMING_STANDARD" else 1,
                    )
                ],
            }
        )
    items.sort(key=lambda item: (0 if item["suggestions"][0]["source"] == "NAMING_STANDARD" else 1, item["item_id"]))
    return {
        "items": items,
        "stats": {
            "actionable_suggestion_rows": len(actionable),
            "unique_items": len(items),
            "deterministic_only_items": deterministic_only,
            "semantic_only_items": semantic_only,
            "overlap_items": overlap,
            "already_decided_excluded": len(decided_items),
            "stale_excluded": len(stale_items),
        },
    }


def _get_reviewable_name_suggestion(db: Session, run_id: uuid.UUID, suggestion_id: uuid.UUID):
    suggestion = db.get(ItemCombinationCleanupNameSuggestion, suggestion_id)
    if suggestion is None or suggestion.run_id != run_id:
        raise CleanupServiceError("Name suggestion not found.", 404)
    if suggestion.status != "PENDING" or suggestion.status == _STALE_NAME_SUGGESTION_STATUS:
        raise CleanupServiceError("Name suggestion is no longer pending.", 409)
    row = get_row(db, run_id, suggestion.item_id_key)
    if row is None:
        raise CleanupServiceError("Cleanup row not found.", 404)
    decision = db.query(ItemCombinationCleanupFieldDecision).filter_by(
        run_id=run_id, item_id_key=row.item_id_key, field_name="name"
    ).first()
    if _name_decision_supersedes_suggestion(decision, suggestion):
        raise CleanupServiceError("A later human Name decision already resolved this suggestion.", 409)
    if (effective_name_for_phase5(db, row) or "") != suggestion.base_value:
        raise CleanupServiceError("Name suggestion baseline is stale.", 409)
    return suggestion, row, decision


def review_name_suggestion(
    db: Session,
    run_id: uuid.UUID,
    suggestion_id: uuid.UUID,
    decision_value: str,
    user: CurrentUser,
    edited_value: str | None = None,
) -> dict:
    suggestion, row, decision = _get_reviewable_name_suggestion(db, run_id, suggestion_id)
    if decision_value == "REJECTED":
        suggestion.status = "REJECTED"
    elif decision_value in {"ACCEPTED", "EDITED"}:
        authoritative = (edited_value or "").strip() if decision_value == "EDITED" else suggestion.suggested_value
        if not authoritative:
            raise CleanupServiceError("An edited Name is required.", 422)
        if decision is None:
            decision = ItemCombinationCleanupFieldDecision(
                run_id=run_id,
                item_id_key=row.item_id_key,
                field_name="name",
                original_value=row.original_item_name,
            )
            db.add(decision)
        decision.deterministic_proposed_value = suggestion.suggested_value
        decision.edited_value = authoritative if decision_value == "EDITED" else None
        decision.decision = "ACCEPTED"
        decision.reviewed_by = uuid.UUID(user.id)
        decision.review_source = f"NAME_SUGGESTION:{suggestion.suggestion_source}"
        decision.reviewed_at = datetime.now(timezone.utc)
        suggestion.status = "ACCEPTED"
    else:
        raise CleanupServiceError("Invalid Name suggestion decision.", 422)
    suggestion.reviewed_by = uuid.UUID(user.id)
    suggestion.reviewed_at = datetime.now(timezone.utc)
    db.flush()
    return {
        "suggestion_id": str(suggestion.id),
        "item_id": suggestion.item_id_key,
        "suggestion_status": suggestion.status,
        "effective_name": effective_name_for_phase5(db, row),
    }


def bulk_accept_name_suggestions(
    db: Session, run_id: uuid.UUID, suggestion_ids: list[uuid.UUID], user: CurrentUser
) -> int:
    if not suggestion_ids:
        raise CleanupServiceError("Select at least one Name suggestion.", 422)
    prepared = []
    for suggestion_id in suggestion_ids:
        suggestion, row, decision = _get_reviewable_name_suggestion(db, run_id, suggestion_id)
        transformations = suggestion.transformations or []
        current = effective_name_for_phase5(db, row)
        replay = suggest_name_standard(current)
        semantic_overlap = db.query(ItemCombinationCleanupNameSuggestion).filter(
            ItemCombinationCleanupNameSuggestion.run_id == run_id,
            ItemCombinationCleanupNameSuggestion.item_id_key == suggestion.item_id_key,
            ItemCombinationCleanupNameSuggestion.suggestion_source == "SEMANTIC_V2",
            ItemCombinationCleanupNameSuggestion.status == "PENDING",
        ).count()
        if (
            suggestion.suggestion_source != "NAMING_STANDARD"
            or len(transformations) != 1
            or transformations[0] not in _BULK_SAFE_NAME_TRANSFORMATIONS
            or replay.suggested_name != suggestion.suggested_value
            or replay.transformations != transformations
            or semantic_overlap
        ):
            raise CleanupServiceError("Bulk Name review is limited to isolated exact approved mappings.", 422)
        prepared.append(suggestion.id)
    for suggestion_id in prepared:
        review_name_suggestion(db, run_id, suggestion_id, "ACCEPTED", user)
    return len(prepared)


def set_field_decision(db: Session, row: ItemCombinationCleanupRow, field_name: str, decision: str, user: CurrentUser, edited_value: str | None = None) -> ItemCombinationCleanupFieldDecision:
    fields = {"name": (row.original_item_name, row.proposed_item_name), "brand": (row.original_brand, row.proposed_brand), "size": (row.original_size, row.proposed_size)}
    if field_name not in fields or fields[field_name][1] is None:
        raise CleanupServiceError("That field has no mutable proposal.", 422)
    if decision not in {"ACCEPTED", "REJECTED", "EDITED"}:
        raise CleanupServiceError("Invalid field decision.", 422)
    if decision == "EDITED" and edited_value is None:
        raise CleanupServiceError("An edited value is required.", 422)
    original, proposed = fields[field_name]
    record = db.query(ItemCombinationCleanupFieldDecision).filter_by(run_id=row.run_id, item_id_key=row.item_id_key, field_name=field_name).first()
    if record is None:
        record = ItemCombinationCleanupFieldDecision(run_id=row.run_id, item_id_key=row.item_id_key, field_name=field_name, original_value=original, deterministic_proposed_value=proposed)
        db.add(record)
    record.decision, record.edited_value, record.reviewed_by, record.review_source, record.reviewed_at = decision, edited_value if decision == "EDITED" else record.edited_value, uuid.UUID(user.id), "UI", datetime.now(timezone.utc)
    return record


def finalize_row_if_all_fields_decided(db: Session, row: ItemCombinationCleanupRow, user: CurrentUser) -> bool:
    """A row leaves the decision inbox only after every real proposal is finalized."""
    proposals = proposed_changes(row)
    if not proposals:
        return False
    decisions = {d.field_name: d.decision for d in db.query(ItemCombinationCleanupFieldDecision).filter_by(run_id=row.run_id, item_id_key=row.item_id_key)}
    if any(decisions.get(field) not in {"ACCEPTED", "REJECTED"} for field in proposals):
        return False
    row.approval_status = "APPROVED" if any(decisions[field] == "ACCEPTED" for field in proposals) else "REJECTED"
    row.reviewed_by, row.reviewed_at = uuid.UUID(user.id), datetime.now(timezone.utc)
    return True


def finalize_individual_field_decisions(db: Session, row: ItemCombinationCleanupRow, user: CurrentUser) -> None:
    if not finalize_row_if_all_fields_decided(db, row, user):
        raise CleanupServiceError("Decide every proposed field before finalizing this row.", 422)


def _sql_text_changed(left, right):
    return func.btrim(func.coalesce(left, "")) != func.btrim(func.coalesce(right, ""))


def _row_out(row: ItemCombinationCleanupRow) -> dict:
    identity = row.source_identity or {}
    item_id = identity.get("item_id")
    if item_id is None:
        item_id = row.item_id_key
    return {
        "item_id": str(item_id),
        "barcode": identity.get("barcode"),
        "mrp": identity.get("mrp"),
        "original_item_name": row.original_item_name,
        "proposed_item_name": row.proposed_item_name,
        "original_brand": row.original_brand,
        "proposed_brand": row.proposed_brand,
        "original_size": row.original_size,
        "proposed_size": row.proposed_size,
        "product_type": row.product_type,
        "classification_confidence": row.classification_confidence,
        "review_status": row.review_status,
        "approval_status": row.approval_status,
        "supplier_name": row.supplier_name,
        "supplier_purchase_date": row.supplier_purchase_date,
        "supplier_purchase_id": row.supplier_purchase_id,
        "supplier_invoice_no": row.supplier_invoice_no,
        "supplier_source_file": row.supplier_source_file,
        "supplier_match_method": row.supplier_match_method,
        "classification_evidence": row.classification_evidence,
        "name_evidence": row.name_evidence,
        "proposed_changes": proposed_changes(row),
        "name_changed": _text_changed(row.original_item_name, row.proposed_item_name),
        "brand_changed": _text_changed(row.original_brand, row.proposed_brand),
        "size_changed": _text_changed(row.original_size, row.proposed_size),
        "row_index": row.row_index,
    }


def semantic_assessment_map(db: Session, run_id: uuid.UUID) -> dict[str, dict]:
    """Latest frozen Phase 4B v2 advisory output, kept outside proposal state."""
    rows = (
        db.query(ItemCombinationCleanupSemanticAssessment)
        .filter(
            ItemCombinationCleanupSemanticAssessment.run_id == run_id,
            ItemCombinationCleanupSemanticAssessment.prompt_version == "phase4b-v2-marker-safety",
            ItemCombinationCleanupSemanticAssessment.validation_status == "VALID",
        )
        .order_by(ItemCombinationCleanupSemanticAssessment.created_at.desc())
        .all()
    )
    assessments: dict[str, dict] = {}
    for assessment in rows:
        if assessment.item_id_key not in assessments:
            assessments[assessment.item_id_key] = {
                "relationship": assessment.relationship,
                "confidence": assessment.confidence,
                "reason": assessment.reason,
                "signals_for_packed": assessment.signals_for_packed or [],
                "signals_against_packed": assessment.signals_against_packed or [],
                "identity_interpretations": assessment.identity_interpretations or [],
            }
    return assessments


def attach_semantic_assessment(row_out: dict, assessment: dict | None) -> dict:
    row_out["semantic_assessment"] = assessment
    row_out["semantic_disagrees_with_packed_proposal"] = bool(
        row_out["product_type"] == "PACKED"
        and assessment
        and assessment["relationship"] == "LIKELY_EXTERNAL_OR_DIFFERENT_PRODUCT"
    )
    return row_out


def attach_field_decisions(db: Session, row_out: dict, run_id: uuid.UUID) -> dict:
    decisions = db.query(ItemCombinationCleanupFieldDecision).filter(ItemCombinationCleanupFieldDecision.run_id == run_id, ItemCombinationCleanupFieldDecision.item_id_key == row_out["item_id"]).all()
    row_out["field_decisions"] = {d.field_name: {"decision": d.decision, "edited_value": d.edited_value} for d in decisions}
    return row_out


def inspection_stats(db: Session, run_id: uuid.UUID) -> dict:
    name_ch = _sql_text_changed(
        ItemCombinationCleanupRow.original_item_name,
        ItemCombinationCleanupRow.proposed_item_name,
    )
    brand_ch = _sql_text_changed(
        ItemCombinationCleanupRow.original_brand,
        ItemCombinationCleanupRow.proposed_brand,
    )
    size_ch = _sql_text_changed(
        ItemCombinationCleanupRow.original_size,
        ItemCombinationCleanupRow.proposed_size,
    )
    change_sum = (
        case((name_ch, 1), else_=0)
        + case((brand_ch, 1), else_=0)
        + case((size_ch, 1), else_=0)
    )
    has_proposed_change = name_ch | brand_ch | size_ch
    safe_mechanical = (
        (ItemCombinationCleanupRow.review_status == "AUTO_APPROVED")
        & (ItemCombinationCleanupRow.product_type != "PACKED")
        & ~brand_ch
        & (name_ch | size_ch)
    )
    batch_reviewable = (
        ItemCombinationCleanupRow.classification_confidence == "HIGH"
    ) & ItemCombinationCleanupRow.product_type.in_(("PURCHASED", "LOOSE")) & brand_ch & (
        ItemCombinationCleanupRow.review_status == "AUTO_APPROVED"
    )
    individual_completed = db.query(func.count()).filter(
        ItemCombinationCleanupRow.run_id == run_id,
        has_proposed_change,
        ~safe_mechanical,
        ~batch_reviewable,
        ItemCombinationCleanupRow.approval_status.in_(("APPROVED", "REJECTED")),
    ).scalar() or 0

    row = db.query(
        func.count().label("total"),
        func.count().filter(ItemCombinationCleanupRow.product_type == "PURCHASED").label("purchased"),
        func.count().filter(ItemCombinationCleanupRow.product_type == "LOOSE").label("loose"),
        func.count().filter(ItemCombinationCleanupRow.product_type == "PACKED").label("packed"),
        func.count().filter(ItemCombinationCleanupRow.product_type == "UNKNOWN").label("unknown"),
        func.count().filter(ItemCombinationCleanupRow.classification_confidence == "HIGH").label("high"),
        func.count().filter(ItemCombinationCleanupRow.classification_confidence == "MEDIUM").label("medium"),
        func.count().filter(ItemCombinationCleanupRow.classification_confidence == "LOW").label("low"),
        func.count().filter(ItemCombinationCleanupRow.review_status == "AUTO_APPROVED").label("auto_approved"),
        func.count().filter(ItemCombinationCleanupRow.review_status == "NEEDS_REVIEW").label("needs_review"),
        func.count().filter(ItemCombinationCleanupRow.review_status == "UNCHANGED").label("unchanged"),
        func.count().filter(name_ch).label("name_changed"),
        func.count().filter(brand_ch).label("brand_changed"),
        func.count().filter(size_ch).label("size_changed"),
        func.count().filter(change_sum > 1).label("multi_field_changed"),
    ).filter(ItemCombinationCleanupRow.run_id == run_id).one()
    return {
        "total": int(row.total or 0),
        "product_type": {
            "PURCHASED": int(row.purchased or 0),
            "LOOSE": int(row.loose or 0),
            "PACKED": int(row.packed or 0),
            "UNKNOWN": int(row.unknown or 0),
        },
        "confidence": {
            "HIGH": int(row.high or 0),
            "MEDIUM": int(row.medium or 0),
            "LOW": int(row.low or 0),
        },
        "review_status": {
            "AUTO_APPROVED": int(row.auto_approved or 0),
            "NEEDS_REVIEW": int(row.needs_review or 0),
            "UNCHANGED": int(row.unchanged or 0),
        },
        "name_changed": int(row.name_changed or 0),
        "brand_changed": int(row.brand_changed or 0),
        "size_changed": int(row.size_changed or 0),
        "multi_field_changed": int(row.multi_field_changed or 0),
        "individual_completed": int(individual_completed),
    }


def _filtered_rows_query(
    db: Session,
    run_id: uuid.UUID,
    *,
    product_type: str | None = None,
    review_status: str | None = None,
    confidence: str | None = None,
    q: str | None = None,
    brand_changed: bool | None = None,
    name_changed: bool | None = None,
    size_changed: bool | None = None,
    approval_status: str | None = None,
    semantic_relationship: str | None = None,
    semantic_confidence: str | None = None,
    semantic_disagreement: bool | None = None,
    semantic_sample: bool = False,
):
    query = db.query(ItemCombinationCleanupRow).filter(
        ItemCombinationCleanupRow.run_id == run_id
    )
    if product_type:
        query = query.filter(ItemCombinationCleanupRow.product_type == product_type)
    if review_status:
        query = query.filter(ItemCombinationCleanupRow.review_status == review_status)
    if confidence:
        query = query.filter(ItemCombinationCleanupRow.classification_confidence == confidence)
    if approval_status:
        query = query.filter(ItemCombinationCleanupRow.approval_status == approval_status)
    if semantic_relationship or semantic_confidence or semantic_disagreement is not None:
        semantic_query = db.query(ItemCombinationCleanupSemanticAssessment.item_id_key).filter(
            ItemCombinationCleanupSemanticAssessment.run_id == run_id,
            ItemCombinationCleanupSemanticAssessment.prompt_version == "phase4b-v2-marker-safety",
            ItemCombinationCleanupSemanticAssessment.validation_status == "VALID",
        )
        if semantic_relationship:
            semantic_query = semantic_query.filter(ItemCombinationCleanupSemanticAssessment.relationship == semantic_relationship)
        if semantic_confidence:
            semantic_query = semantic_query.filter(ItemCombinationCleanupSemanticAssessment.confidence == semantic_confidence)
        if semantic_disagreement is not None:
            query = query.filter(ItemCombinationCleanupRow.product_type == "PACKED")
            relation = "LIKELY_EXTERNAL_OR_DIFFERENT_PRODUCT"
            semantic_query = semantic_query.filter(
                ItemCombinationCleanupSemanticAssessment.relationship == relation if semantic_disagreement
                else ItemCombinationCleanupSemanticAssessment.relationship != relation
            )
        query = query.filter(ItemCombinationCleanupRow.item_id_key.in_(semantic_query))
    if semantic_sample:
        query = query.filter(ItemCombinationCleanupRow.item_id_key.in_(semantic_sample_item_ids(db, run_id)))
    if brand_changed is not None:
        expr = _sql_text_changed(
            ItemCombinationCleanupRow.original_brand,
            ItemCombinationCleanupRow.proposed_brand,
        )
        query = query.filter(expr if brand_changed else ~expr)
    if name_changed is not None:
        expr = _sql_text_changed(
            ItemCombinationCleanupRow.original_item_name,
            ItemCombinationCleanupRow.proposed_item_name,
        )
        query = query.filter(expr if name_changed else ~expr)
    if size_changed is not None:
        expr = _sql_text_changed(
            ItemCombinationCleanupRow.original_size,
            ItemCombinationCleanupRow.proposed_size,
        )
        query = query.filter(expr if size_changed else ~expr)
    needle = (q or "").strip()
    if needle:
        like = f"%{needle}%"
        barcode_text = ItemCombinationCleanupRow.source_identity["barcode"].astext
        query = query.filter(
            or_(
                ItemCombinationCleanupRow.item_id_key.ilike(like),
                ItemCombinationCleanupRow.original_item_name.ilike(like),
                ItemCombinationCleanupRow.proposed_item_name.ilike(like),
                ItemCombinationCleanupRow.original_brand.ilike(like),
                ItemCombinationCleanupRow.proposed_brand.ilike(like),
                ItemCombinationCleanupRow.supplier_name.ilike(like),
                barcode_text.ilike(like),
            )
        )
    return query


def _semantic_sample_key(run_id: uuid.UUID, row: ItemCombinationCleanupRow, assessment: dict) -> tuple:
    """Stable diversity bucket; a review convenience, never classification data."""
    evidence = row.classification_evidence or {}
    identity = row.source_identity or {}
    barcode = str(identity.get("barcode") or "")
    barcode_form = "standard" if barcode.isdigit() and len(barcode) in (8, 12, 13, 14) else "internal"
    words = " ".join((row.original_item_name or "").split()[:2]).upper()
    bucket = (
        assessment.get("confidence"), row.classification_confidence,
        bool(evidence.get("has_phm_marker")), bool(evidence.get("external_brand_candidates")),
        barcode_form, words,
    )
    digest = hashlib.sha256(f"{run_id}:{row.item_id_key}".encode()).hexdigest()
    return bucket + (digest,)


def semantic_sample_item_ids(db: Session, run_id: uuid.UUID, limit: int = 25) -> list[str]:
    """Round-robin deterministic diversity sample of AI-supported PACKED rows."""
    assessments = semantic_assessment_map(db, run_id)
    rows = db.query(ItemCombinationCleanupRow).filter(
        ItemCombinationCleanupRow.run_id == run_id,
        ItemCombinationCleanupRow.product_type == "PACKED",
    ).all()
    buckets: dict[tuple, list[ItemCombinationCleanupRow]] = defaultdict(list)
    for row in rows:
        assessment = assessments.get(row.item_id_key)
        if assessment and assessment.get("relationship") == "LIKELY_PACKED_VERSION":
            key = _semantic_sample_key(run_id, row, assessment)[:-1]
            buckets[key].append(row)
    for key, bucket_rows in buckets.items():
        bucket_rows.sort(key=lambda row: _semantic_sample_key(run_id, row, assessments[row.item_id_key])[-1])
    selected: list[str] = []
    ordered = sorted(buckets, key=lambda key: hashlib.sha256(repr(key).encode()).hexdigest())
    while ordered and len(selected) < limit:
        next_round = []
        for key in ordered:
            if buckets[key] and len(selected) < limit:
                selected.append(buckets[key].pop(0).item_id_key)
            if buckets[key]:
                next_round.append(key)
        ordered = next_round
    return selected


def list_rows(
    db: Session,
    run_id: uuid.UUID,
    *,
    product_type: str | None = None,
    review_status: str | None = None,
    confidence: str | None = None,
    q: str | None = None,
    brand_changed: bool | None = None,
    name_changed: bool | None = None,
    size_changed: bool | None = None,
    approval_status: str | None = None,
    semantic_relationship: str | None = None,
    semantic_confidence: str | None = None,
    semantic_disagreement: bool | None = None,
    semantic_sample: bool = False,
    limit: int = 50,
    offset: int = 0,
) -> tuple[int, list[ItemCombinationCleanupRow], dict]:
    query = _filtered_rows_query(
        db,
        run_id,
        product_type=product_type,
        review_status=review_status,
        confidence=confidence,
        q=q,
        brand_changed=brand_changed,
        name_changed=name_changed,
        size_changed=size_changed,
        approval_status=approval_status,
        semantic_relationship=semantic_relationship,
        semantic_confidence=semantic_confidence,
        semantic_disagreement=semantic_disagreement,
        semantic_sample=semantic_sample,
    )
    total = query.count()
    rows = (
        query.order_by(ItemCombinationCleanupRow.row_index)
        .offset(offset)
        .limit(limit)
        .all()
    )
    return total, rows, inspection_stats(db, run_id)


def _mechanical_only_auto_approval(row: ItemCombinationCleanupRow) -> bool:
    return (
        row.review_status == "AUTO_APPROVED"
        and row.product_type != "PACKED"
        and not _text_changed(row.original_brand, row.proposed_brand)
        and (_text_changed(row.original_item_name, row.proposed_item_name)
             or _text_changed(row.original_size, row.proposed_size))
    )


def is_export_eligible(row: ItemCombinationCleanupRow) -> bool:
    return row.approval_status == "APPROVED" or _mechanical_only_auto_approval(row)


def set_approval(row: ItemCombinationCleanupRow, status: str, user: CurrentUser) -> ItemCombinationCleanupRow:
    row.approval_status = status
    row.reviewed_by = uuid.UUID(user.id)
    row.reviewed_at = datetime.now(timezone.utc)
    return row


def bulk_approve(db: Session, run_id: uuid.UUID, filters: dict, user: CurrentUser) -> int:
    meaningful = any(filters.get(key) not in (None, "") for key in (
        "product_type", "review_status", "confidence", "q", "brand_changed", "name_changed", "size_changed"
    ))
    if not meaningful:
        raise CleanupServiceError("Bulk approval requires at least one filter.", 422)
    query = _filtered_rows_query(db, run_id, **filters).filter(
        ItemCombinationCleanupRow.approval_status == "PENDING"
    )
    if query.filter(ItemCombinationCleanupRow.classification_confidence == "LOW").count():
        raise CleanupServiceError("LOW confidence rows cannot be bulk-approved.", 422)
    rows = query.all()
    for row in rows:
        set_approval(row, "APPROVED", user)
    db.flush()
    return len(rows)


def bulk_decide(db: Session, run_id: uuid.UUID, status: str, user: CurrentUser, item_ids: list[str], filters: dict) -> tuple[int, dict[str, int]]:
    query = db.query(ItemCombinationCleanupRow).filter(ItemCombinationCleanupRow.run_id == run_id, ItemCombinationCleanupRow.approval_status == "PENDING")
    if item_ids:
        query = query.filter(ItemCombinationCleanupRow.item_id_key.in_(item_ids))
    else:
        if not any(filters.values()):
            raise CleanupServiceError("Select rows or provide a queue filter.", 422)
        if filters.get("type"): query = query.filter(ItemCombinationCleanupRow.product_type == filters["type"])
        if filters.get("confidence"): query = query.filter(ItemCombinationCleanupRow.classification_confidence == filters["confidence"])
        if filters.get("supplier_match_method"): query = query.filter(ItemCombinationCleanupRow.supplier_match_method == filters["supplier_match_method"])
        if filters.get("semantic_relationship") or filters.get("semantic_confidence") or filters.get("semantic_disagreement") is not None:
            semantic_query = db.query(ItemCombinationCleanupSemanticAssessment.item_id_key).filter(
                ItemCombinationCleanupSemanticAssessment.run_id == run_id,
                ItemCombinationCleanupSemanticAssessment.prompt_version == "phase4b-v2-marker-safety",
                ItemCombinationCleanupSemanticAssessment.validation_status == "VALID",
            )
            if filters.get("semantic_relationship"):
                semantic_query = semantic_query.filter(ItemCombinationCleanupSemanticAssessment.relationship == filters["semantic_relationship"])
            if filters.get("semantic_confidence"):
                semantic_query = semantic_query.filter(ItemCombinationCleanupSemanticAssessment.confidence == filters["semantic_confidence"])
            if filters.get("semantic_disagreement") is not None:
                query = query.filter(ItemCombinationCleanupRow.product_type == "PACKED")
                relation = "LIKELY_EXTERNAL_OR_DIFFERENT_PRODUCT"
                semantic_query = semantic_query.filter(ItemCombinationCleanupSemanticAssessment.relationship == relation if filters["semantic_disagreement"] else ItemCombinationCleanupSemanticAssessment.relationship != relation)
            query = query.filter(ItemCombinationCleanupRow.item_id_key.in_(semantic_query))
    rows = query.all()
    from collections import Counter
    summary = Counter(row.product_type for row in rows)
    for row in rows: set_approval(row, status, user)
    db.flush()
    return len(rows), dict(summary)


def ai_supported_packed_pending_count(db: Session, run_id: uuid.UUID) -> int:
    query = db.query(ItemCombinationCleanupRow).filter(
        ItemCombinationCleanupRow.run_id == run_id,
        ItemCombinationCleanupRow.product_type == "PACKED",
        ItemCombinationCleanupRow.approval_status == "PENDING",
    )
    semantic_rows = db.query(ItemCombinationCleanupSemanticAssessment.item_id_key).filter(
        ItemCombinationCleanupSemanticAssessment.run_id == run_id,
        ItemCombinationCleanupSemanticAssessment.prompt_version == "phase4b-v2-marker-safety",
        ItemCombinationCleanupSemanticAssessment.validation_status == "VALID",
        ItemCombinationCleanupSemanticAssessment.relationship == "LIKELY_PACKED_VERSION",
    )
    return query.filter(ItemCombinationCleanupRow.item_id_key.in_(semantic_rows)).count()


def export_approved_run(db: Session, run: ItemCombinationCleanupRun) -> ItemCombinationCleanupRun:
    source = Path(run.original_path)
    if not source.is_file():
        raise CleanupServiceError("Original workbook is missing on disk.", 404)
    proposals: dict[str, dict[str, object]] = {}
    decisions = {(d.item_id_key, d.field_name): d for d in db.query(ItemCombinationCleanupFieldDecision).filter(ItemCombinationCleanupFieldDecision.run_id == run.id)}
    for row in db.query(ItemCombinationCleanupRow).filter(ItemCombinationCleanupRow.run_id == run.id):
        changes = {}
        for field, source_column, original, proposed in (("name", "Item Name", row.original_item_name, row.proposed_item_name), ("brand", "Brand", row.original_brand, row.proposed_brand), ("size", "Size", row.original_size, row.proposed_size)):
            record = decisions.get((row.item_id_key, field))
            if record:
                reviewed_proposal = record.deterministic_proposed_value
                value = resolved_reviewed_value(original, reviewed_proposal, record.decision, record.edited_value)
                if _text_changed(original, value): changes[source_column] = value
            elif is_export_eligible(row) and _text_changed(original, proposed):
                changes[source_column] = proposed
        if changes: proposals[row.item_id_key] = changes
    dest = run_dir(run.id) / "approved_output.xlsx"
    try:
        export_approved_proposals(source, dest, proposals)
    except WorkbookValidationError as exc:
        run.approved_output_path = None
        raise CleanupServiceError(str(exc), 422) from exc
    run.approved_output_path = str(dest)
    db.flush()
    return run


def get_row(
    db: Session, run_id: uuid.UUID, item_id: str
) -> ItemCombinationCleanupRow | None:
    return (
        db.query(ItemCombinationCleanupRow)
        .filter(
            ItemCombinationCleanupRow.run_id == run_id,
            ItemCombinationCleanupRow.item_id_key == item_id,
        )
        .first()
    )


def export_proposals_audit(db: Session, run: ItemCombinationCleanupRun) -> Path:
    """Write a read-only audit workbook. Does not touch source.xlsx or ER4U export."""
    import json

    from openpyxl import Workbook

    rows = (
        db.query(ItemCombinationCleanupRow)
        .filter(ItemCombinationCleanupRow.run_id == run.id)
        .order_by(ItemCombinationCleanupRow.row_index)
        .all()
    )
    dest = run_dir(run.id) / "proposals_audit.xlsx"
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Proposal_Audit"
    ws.append([
        "Item_Id",
        "Barcode",
        "MRP",
        "Original Item Name",
        "Proposed Item Name",
        "Original Brand",
        "Proposed Brand",
        "Original Size",
        "Proposed Size",
        "Product Type",
        "Confidence",
        "Review Status",
        "Supplier",
        "Purchase Date",
        "Purchase ID",
        "Invoice No",
        "Source File",
        "Match Method",
        "Name Evidence",
        "Classification Evidence",
    ])
    for row in rows:
        out = _row_out(row)
        ws.append([
            out["item_id"],
            out["barcode"],
            out["mrp"],
            out["original_item_name"],
            out["proposed_item_name"],
            out["original_brand"],
            out["proposed_brand"],
            out["original_size"],
            out["proposed_size"],
            out["product_type"],
            out["classification_confidence"],
            out["review_status"],
            out["supplier_name"],
            out["supplier_purchase_date"].isoformat() if out["supplier_purchase_date"] else None,
            out["supplier_purchase_id"],
            out["supplier_invoice_no"],
            out["supplier_source_file"],
            out["supplier_match_method"],
            json.dumps(out["name_evidence"], ensure_ascii=False) if out["name_evidence"] is not None else None,
            json.dumps(out["classification_evidence"], ensure_ascii=False)
            if out["classification_evidence"] is not None else None,
        ])
    wb.save(dest)
    wb.close()
    return dest


def export_staff_review(db: Session, run: ItemCombinationCleanupRun, item_ids: list[str], packed_only: bool) -> Path:
    """Human-review workbook; never an ER4U output and never mutates the source."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    rows = db.query(ItemCombinationCleanupRow).filter(ItemCombinationCleanupRow.run_id == run.id, ItemCombinationCleanupRow.review_status == "NEEDS_REVIEW").order_by(ItemCombinationCleanupRow.row_index)
    if item_ids: rows = rows.filter(ItemCombinationCleanupRow.item_id_key.in_(item_ids))
    if packed_only: rows = rows.filter(ItemCombinationCleanupRow.product_type == "PACKED")
    semantic = semantic_assessment_map(db, run.id)
    wb = Workbook(); ins = wb.active; ins.title = "Instructions"; ws = wb.create_sheet("Review Items")
    ins.sheet_view.showGridLines = False
    ins["A1"] = "AxonFlux staff review"; ins["A1"].font = Font(bold=True, size=16)
    for i, text in enumerate(["This workbook contains Item Combination changes that AxonFlux could not safely finalize automatically.", "For PACKED products answer: Do we pack this item ourselves? YES, NO, or NOT SURE.", "For supplier review answer: Is the proposed supplier correct? YES, NO, or NOT SURE.", "Do not change Item_Id or Barcode. Only yellow cells are intended for staff input.", "Blue = AI supports PACKED; pink = AI flags external/different; yellow = uncertain or staff input; gray = read-only reference."], 3): ins.cell(i, 1, text)
    headers = ["Review ID", "Item_Id", "Current Item Name", "Proposed Item Name", "Current Brand", "Proposed Brand", "Current Size", "Proposed Size", "Product Type", "Why AxonFlux flagged this", "AI Assessment", "AI Confidence", "Matched LOOSE Product", "Latest Supplier", "Staff Decision", "Correct Item Name", "Correct Brand", "Correct Size", "Staff Notes"]
    ws.append(headers); ws.freeze_panes = "A2"; ws.auto_filter.ref = "A1:S1"; ws.sheet_view.showGridLines = False
    gray, yellow, blue, pink = PatternFill("solid", fgColor="E7E6E6"), PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="FCE4D6")
    for cell in ws[1]: cell.font = Font(bold=True); cell.fill = gray
    for index, row in enumerate(rows, 1):
        evidence = row.classification_evidence or {}; assessment = semantic.get(row.item_id_key) or {}; identity = row.source_identity or {}
        question = "Do we pack this item ourselves?" if row.product_type == "PACKED" else "Is the proposed supplier correct?"
        ws.append([f"{run.id}:{row.item_id_key}", row.item_id_key, row.original_item_name, row.proposed_item_name, row.original_brand, row.proposed_brand, row.original_size, row.proposed_size, row.product_type, question, assessment.get("relationship"), assessment.get("confidence"), evidence.get("source_loose_item_name"), row.supplier_name, None, None, None, None, None])
        excel_row = index + 1
        for col in range(1, 15): ws.cell(excel_row, col).fill = gray; ws.cell(excel_row, col).protection = Protection(locked=True)
        for col in range(15, 20): ws.cell(excel_row, col).fill = yellow; ws.cell(excel_row, col).protection = Protection(locked=False)
        fill = blue if assessment.get("relationship") == "LIKELY_PACKED_VERSION" else pink if assessment.get("relationship") == "LIKELY_EXTERNAL_OR_DIFFERENT_PRODUCT" else yellow
        ws.cell(excel_row, 11).fill = fill
    validation = DataValidation(type="list", formula1='"YES,NO,NOT SURE"'); ws.add_data_validation(validation); validation.add(f"O2:O{ws.max_row}")
    for column, width in {"A": 40, "B": 12, "C": 30, "D": 30, "E": 22, "F": 22, "G": 14, "H": 14, "I": 14, "J": 34, "K": 28, "L": 14, "M": 30, "N": 26, "O": 18, "P": 30, "Q": 28, "R": 16, "S": 34}.items(): ws.column_dimensions[column].width = width
    for row in ws.iter_rows():
        for cell in row: cell.alignment = __import__("openpyxl").styles.Alignment(wrap_text=True, vertical="top")
    dest = run_dir(run.id) / "staff_review.xlsx"; wb.save(dest); return dest


def _staff_review_plan(db: Session, run: ItemCombinationCleanupRun, content: bytes) -> dict:
    """Parse a staff workbook into explicit, non-mutating field actions."""
    from io import BytesIO
    from openpyxl import load_workbook
    try: ws = load_workbook(BytesIO(content), read_only=True, data_only=True)["Review Items"]
    except Exception as exc: raise CleanupServiceError("Invalid staff review workbook.", 422) from exc
    values = list(ws.values)
    if not values: raise CleanupServiceError("Staff review workbook is empty.", 422)
    headers = {str(name): i for i, name in enumerate(values[0])}
    required = {"Review ID", "Item_Id", "Staff Decision"}
    if not required.issubset(headers): raise CleanupServiceError("Staff review workbook has missing required columns.", 422)
    rows = {r.item_id_key: r for r in db.query(ItemCombinationCleanupRow).filter(ItemCombinationCleanupRow.run_id == run.id)}
    decisions = {(d.item_id_key, d.field_name): d for d in db.query(ItemCombinationCleanupFieldDecision).filter(ItemCombinationCleanupFieldDecision.run_id == run.id)}
    seen, planned, summary = set(), [], {"applicable": 0, "protected": 0, "not_sure": 0, "invalid": 0}
    for raw in values[1:]:
        cell = lambda header: raw[headers[header]] if header in headers and headers[header] < len(raw) else None
        item_id = str(cell("Item_Id") or "").strip(); review_id = str(cell("Review ID") or "")
        if not item_id: continue
        if item_id in seen: raise CleanupServiceError(f"Duplicate Item_Id: {item_id}", 422)
        seen.add(item_id)
        if item_id not in rows: raise CleanupServiceError(f"Unknown Item_Id: {item_id}", 422)
        if not review_id.startswith(f"{run.id}:{item_id}"): raise CleanupServiceError(f"Review ID does not match Item_Id: {item_id}", 422)
        decision = str(cell("Staff Decision") or "").strip().upper()
        if decision not in {"YES", "NO", "", "NOT SURE"}: raise CleanupServiceError(f"Invalid staff decision for Item_Id: {item_id}", 422)
        row = rows[item_id]
        if row.approval_status != "PENDING":
            summary["protected"] += 1; planned.append({"item_id": item_id, "eligibility": "PROTECTED", "reason": "Row already resolved", "actions": []}); continue
        if decision in {"", "NOT SURE"}:
            summary["not_sure"] += 1; planned.append({"item_id": item_id, "eligibility": "NOT_SURE", "actions": []}); continue
        changes = proposed_changes(row); actions = []
        for field in ("name", "brand", "size"):
            if field not in changes or (item_id, field) in decisions: continue
            correct_header = {"name": "Correct Item Name", "brand": "Correct Brand", "size": "Correct Size"}[field]
            corrected = cell(correct_header)
            corrected = None if corrected is None or str(corrected).strip() == "" else str(corrected).strip()
            if corrected is not None: action, value = "ACCEPT", corrected
            elif decision == "YES": action, value = "ACCEPT", changes[field]["proposed"]
            elif field == "brand": action, value = "REJECT", changes[field]["original"]
            else: continue
            actions.append({"field": field, "action": action, "from": changes[field]["original"], "to": value, "corrected": corrected is not None})
        eligibility = "APPLICABLE" if actions else "PROTECTED"
        summary["applicable" if actions else "protected"] += 1
        planned.append({"item_id": item_id, "eligibility": eligibility, "staff_decision": decision, "actions": actions})
    return {"file_hash": hashlib.sha256(content).hexdigest(), "summary": summary, "rows": planned}


def preview_staff_review_import(db: Session, run: ItemCombinationCleanupRun, content: bytes) -> dict:
    result = _staff_review_plan(db, run, content)
    result["preview_only"] = True
    return result


def apply_staff_review_import(db: Session, run: ItemCombinationCleanupRun, content: bytes, preview_hash: str, user: CurrentUser) -> dict:
    if hashlib.sha256(content).hexdigest() != preview_hash: raise CleanupServiceError("Workbook changed after preview.", 409)
    plan = _staff_review_plan(db, run, content); applied_rows = applied_fields = 0
    for item in plan["rows"]:
        if item["eligibility"] != "APPLICABLE": continue
        row = db.query(ItemCombinationCleanupRow).filter_by(run_id=run.id, item_id_key=item["item_id"]).with_for_update().one()
        if row.approval_status != "PENDING": raise CleanupServiceError(f"Item {row.item_id_key} changed after preview.", 409)
        for action in item["actions"]:
            if db.query(ItemCombinationCleanupFieldDecision).filter_by(run_id=run.id, item_id_key=row.item_id_key, field_name=action["field"]).first(): raise CleanupServiceError(f"Item {row.item_id_key} changed after preview.", 409)
            original, proposed = {"name": (row.original_item_name, row.proposed_item_name), "brand": (row.original_brand, row.proposed_brand), "size": (row.original_size, row.proposed_size)}[action["field"]]
            db.add(ItemCombinationCleanupFieldDecision(run_id=run.id, item_id_key=row.item_id_key, field_name=action["field"], original_value=original, deterministic_proposed_value=proposed, edited_value=action["to"] if action["corrected"] else None, decision="ACCEPTED" if action["action"] == "ACCEPT" else "REJECTED", reviewed_by=uuid.UUID(user.id), review_source="STAFF_IMPORT", reviewed_at=datetime.now(timezone.utc)))
            applied_fields += 1
        db.flush(); finalize_row_if_all_fields_decided(db, row, user); applied_rows += 1
    return {"applied_rows": applied_rows, "applied_fields": applied_fields, "protected_rows": plan["summary"]["protected"], "not_sure_rows": plan["summary"]["not_sure"]}
