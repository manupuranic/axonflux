from openpyxl import Workbook,load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from config.db import engine
from sqlalchemy import text
from api.tools.item_combination_cleanup.name_semantic import detect,context_qualified
RID='49a0b23a-9cc9-4b9c-b487-e80bb4d76f48'; OUT='data/phase5b_name_context_evaluation.xlsx'
r=list(engine.connect().execute(text("select item_id_key,original_item_name,source_identity->>'barcode' from app.item_combination_cleanup_rows where run_id=:r"),{'r':RID})); names=[(str(i),n or '') for i,n,_ in r]; bm={str(b):str(i) for i,_,b in r}; p={}
for b,n in engine.connect().execute(text('select barcode,item_name_raw from raw.raw_purchase_itemwise where item_name_raw is not null')):
 if str(b) in bm:p.setdefault(bm[str(b)],[]).append(n)
s=context_qualified(detect(names),p,[n for _,n in names]); wanted=['CHOCOLT','SHOULDE','CHESE','CHIPPS','WIPERC','FANTACY','FOXTAL','LIQUD']; hard=['LIVON','SREE','DUAL','MEAT','KANJI','TONG','CRY','DIA']; chosen=[]; used=set()
for x in s:
 if x['token'] in wanted+hard and x['item_id'] not in used: chosen.append(x);used.add(x['item_id'])
for x in s:
 if len(chosen)>=40:break
 if x['item_id'] not in used and x['token'] not in {z['token'] for z in chosen}:chosen.append(x);used.add(x['item_id'])
for i,n in names:
 if len(chosen)>=58:break
 if i not in used and any(t in n.upper() for t in ['JAWARI','AVALAKKI','SAJJE','PEANUT','CREM','RINIVA']):chosen.append({'item_id':i,'name':n,'token':'','evidence':'','reason':'NEGATIVE_CONTROL','context_evidence':['PROTECTED']});used.add(i)
wb=Workbook();ws=wb.active;ws.title='Evaluation';ws.append(['Item_Id','Effective Item Name','Suspicious token','Proposed lexical neighbor','Detector reason','Corroborating evidence','Human Label','Human Corrected Name','Human Notes'])
for x in chosen:ws.append([x['item_id'],x['name'],x['token'],x['evidence'],x['reason'],', '.join(x.get('context_evidence',[])),'','',''])
dv=DataValidation(type='list',formula1='"CORRECTION,NO_CHANGE,UNCERTAIN"');ws.add_data_validation(dv);dv.add(f'G2:G{ws.max_row}');ins=wb.create_sheet('Instructions');ins['A1']='Use CORRECTION only for genuine naming errors; use NO_CHANGE for legitimate brands, local terms, models, or variants; use UNCERTAIN when evidence is insufficient.';wb.save(OUT)
w=load_workbook(OUT);assert w.sheetnames==['Evaluation','Instructions'] and w['Evaluation'].data_validations.count==1 and all(w['Evaluation'].cell(i,7).value is None for i in range(2,w['Evaluation'].max_row+1));print(w['Evaluation'].max_row-1)
