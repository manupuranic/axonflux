from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from config.db import engine
from sqlalchemy import text
from api.tools.item_combination_cleanup.name_semantic import detect, tokens

RID="49a0b23a-9cc9-4b9c-b487-e80bb4d76f48"
OUT=Path("data/phase5b_name_evaluation.xlsx")
rows=list(engine.connect().execute(text("select item_id_key,original_item_name from app.item_combination_cleanup_rows where run_id=:r"),{"r":RID}))
names=[(str(i),n or "") for i,n in rows]; candidates=detect(names)
trunc=[x for x in candidates if x["reason"]=="TRUNCATED_LOOKING"][:28]
used={x["item_id"] for x in trunc}; rare=[x for x in candidates if x["reason"]=="RARE_NEAR_NEIGHBOR" and x["item_id"] not in used][:32]; used|={x["item_id"] for x in rare}
negative=[]
for i,n in names:
    if i not in used and any(t in n.upper() for t in ("JAWARI","AVALAKKI","PEANUT","CREM","RINIVA","PARTH","ANTAJI")):
        negative.append({"item_id":i,"name":n,"token":"","reason":"NEGATIVE_CONTROL","evidence":"Protected/local/unusual terminology"}); used.add(i)
    if len(negative)==18: break
sample=trunc+rare+negative
freq=Counter(t for _,n in names for t in tokens(n))
wb=Workbook(); ws=wb.active; ws.title="Evaluation"
ws.append(["Item_Id","Effective Item Name","Suspicious token/phrase","Detector reason","Nearest vocabulary candidate","Similarity/distance evidence","Token/catalog frequency","Relevant purchase-name evidence","Human Label","Human Corrected Name","Human Notes"])
for x in sample:
    ws.append([x["item_id"],x["name"],x["token"],x["reason"],x.get("evidence",""),x.get("evidence",""),freq.get(x["token"],0),"","","",""])
dv=DataValidation(type="list",formula1='"CORRECTION,NO_CHANGE,UNCERTAIN"'); ws.add_data_validation(dv); dv.add(f"I2:I{ws.max_row}")
ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
for c,w in zip("ABCDEFGHIJK",[12,42,22,24,30,30,18,35,18,42,32]): ws.column_dimensions[c].width=w
ins=wb.create_sheet("Instructions"); ins["A1"]="CORRECTION only for a genuine naming error. Legitimate unusual brand, local, product, or model terminology is NO_CHANGE. Use UNCERTAIN when evidence is insufficient."
wb.save(OUT)
check=load_workbook(OUT); assert check.sheetnames==["Evaluation","Instructions"]; assert check["Evaluation"].max_row==len(sample)+1; assert check["Evaluation"].data_validations.count==1; assert all(check["Evaluation"].cell(r,9).value is None for r in range(2,check["Evaluation"].max_row+1)); print({"rows":len(sample),"truncated":len(trunc),"rare":len(rare),"negative":len(negative),"path":str(OUT)})
