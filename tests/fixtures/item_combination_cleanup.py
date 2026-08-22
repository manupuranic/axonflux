"""Helpers for Phase 1 Item Combination Cleanup tests."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADERS = [
    "Item_Id",
    "ItemStatus",
    "Barcode",
    "Item Name",
    "HsnCode",
    "Tax Category",
    "Brand",
    "Size",
    "Colour",
    "Style",
    "Expiry Date",
    "Min Stock",
    "Max Stock",
    "Purchase Price",
    "MRP",
    "Rate",
    "System Stock",
    "ExtraNote",
]


def write_fixture_workbook(path: Path) -> Path:
    """Synthetic master: duplicate barcodes, blank Brand/Size, extra column, string Item_Id."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Item_Master_List"
    ws.append(HEADERS)
    ws.append([
        "1001", 1, "8900000000001", "DABUR RED TOOTHPASTE 100GM", "33051090",
        "GST(18)", "DAB", "100G", None, None, None, 0, 0, 40, 55, 50, 10, "keep-me",
    ])
    ws.append([
        "1002", 1, "8900000000001", "DABUR RED TOOTHPASTE 100GM", "33051090",
        "GST(18)", None, None, None, None, None, 0, 0, 38, 49, 48, 4, "same-barcode-diff-mrp",
    ])
    ws.append([
        "1003", 1, "AH61", "SAME LOOSE", "10082100",
        "GST(5)", None, None, None, None, None, 0, 0, 115, 175, 155, 20, None,
    ])
    ws.append([
        "1004", 1, "AP96", "SAME 1KG", "10082100",
        "GST(5)", "PKG", "1KG", None, None, None, 0, 0, None, 175, 155, 8, None,
    ])
    ws["A6"] = "001234"
    ws["B6"] = 1
    ws["C6"] = "BC001"
    ws["D6"] = "LEADING ZERO ID"
    ws["E6"] = "21069099"
    ws["F6"] = "GST(18)"
    ws["G6"] = None
    ws["H6"] = None
    ws["I6"] = None
    ws["J6"] = None
    ws["K6"] = None
    ws["L6"] = 0
    ws["M6"] = 0
    ws["N6"] = 10
    ws["O6"] = 20
    ws["P6"] = 18
    ws["Q6"] = 1
    ws["R6"] = "string-id"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


PHASE2_ROWS = [
    ["P2-1001", 1, "ICC2_SAME_LOOSE", "SAME LOOSE", "10082100",
     "GST(5)", None, None, None, None, None, 0, 0, 115, 175, 155, 20, None],
    ["P2-1002", 1, "ICC2_SAME_1KG", "SAME 1KG", "10082100",
     "GST(5)", "PKG", None, None, None, None, 0, 0, None, 175, 155, 8, None],
    ["P2-1003", 1, "ICC2_TEA", "green tea 100gm", "09021000",
     "GST(5)", None, None, None, None, None, 0, 0, 80, 120, 110, 3, None],
    ["P2-1004", 1, "ICC2_PKG_BUY", "SUNFLOWER OIL 1LTR", "15121900",
     "GST(5)", "PKG", "1LTR", None, None, None, 0, 0, 90, 140, 130, 6, None],
    ["P2-1005", 1, "ICC2_PHM", "WHEAT FLOUR (PHM) 1KG", "11010000",
     "GST(5)", None, None, None, None, None, 0, 0, 40, 60, 55, 12, None],
    ["P2-1006", 1, "ICC2_ADU", "ADUKALE NIPPATTU", "19059090",
     "GST(18)", None, None, None, None, None, 0, 0, 50, 80, 75, 4, None],
    ["P2-1007", 1, "ICC2_LOOSE_NONE", "RICE LOOSE", "10063000",
     "GST(5)", "KEEPME", None, None, None, None, None, 0, 0, 30, 45, 40, 9, None],
    ["P2-1008", 1, "ICC2_SHARED", "SOAP A 100G", "34011100",
     "GST(18)", None, None, None, None, None, 0, 0, 10, 25, 22, 2, None],
    ["P2-1009", 1, "ICC2_SHARED", "SOAP B 100G", "34011100",
     "GST(18)", None, None, None, None, None, 0, 0, 10, 25, 22, 1, None],
    ["P2-1010", 1, "ICC2_FALLBACK", "BISCUIT PACK 50G", "19053100",
     "GST(18)", None, None, None, None, None, 0, 0, 8, 99, 90, 5, None],
]


def write_phase2_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Item_Master_List"
    ws.append(HEADERS)
    for row in PHASE2_ROWS:
        ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


PHASE2_PURCHASES = [
    {
        "barcode": "ICC2_SAME_LOOSE",
        "mrp": 175,
        "supplier": "SRI BALAJI TRADERS",
        "purchase_date": "12/07/2026",
        "purchase_id": "PO-LOOSE-1",
        "invoice_no": "INV-LOOSE-1",
    },
    {
        "barcode": "ICC2_PKG_BUY",
        "mrp": 140,
        "supplier": "OIL DISTRIBUTORS",
        "purchase_date": "01-06-2026",
        "purchase_id": "PO-OIL-1",
        "invoice_no": "INV-OIL-1",
    },
    {
        "barcode": "ICC2_PHM",
        "mrp": 60,
        "supplier": "EXTERNAL MILLERS",
        "purchase_date": "15/05/2026",
        "purchase_id": "PO-PHM-1",
        "invoice_no": "INV-PHM-1",
    },
    {
        "barcode": "ICC2_ADU",
        "mrp": 80,
        "supplier": "ADUKALE FOODS",
        "purchase_date": "20-04-2026",
        "purchase_id": "PO-ADU-1",
        "invoice_no": "INV-ADU-1",
    },
    {
        "barcode": "ICC2_SHARED",
        "mrp": 25,
        "supplier": "SOAP WHOLESALE",
        "purchase_date": "10/03/2026",
        "purchase_id": "PO-SOAP-1",
        "invoice_no": "INV-SOAP-1",
    },
    {
        "barcode": "ICC2_FALLBACK",
        "mrp": 10,
        "supplier": "SNACK SUPPLIER",
        "purchase_date": "05-02-2026",
        "purchase_id": "PO-FB-1",
        "invoice_no": "INV-FB-1",
    },
]


def mutate_column(path: Path, header: str, row: int, value) -> None:
    wb = load_workbook(path)
    ws = wb.active
    col = None
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == header:
            col = idx
            break
    assert col is not None, f"header {header} not found"
    ws.cell(row, col).value = value
    wb.save(path)
    wb.close()
