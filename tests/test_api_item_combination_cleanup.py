"""API tests for Item Combination Cleanup Phase 1 (upload → export → download)."""

import uuid
from types import SimpleNamespace

from sqlalchemy import text

from tests.conftest import test_engine
from tests.fixtures.item_combination_cleanup import (
    PHASE2_PURCHASES,
    write_fixture_workbook,
    write_phase2_workbook,
)


def test_proposed_changes_treats_null_as_absent_and_empty_string_as_clear():
    from api.tools.item_combination_cleanup.service import proposed_changes

    absent = SimpleNamespace(
        original_item_name="COCONUT SCRUB PAD (PANAI)", proposed_item_name="COCONUT SCRUB PAD (PANAI)",
        original_brand="Panai", proposed_brand="Panai", original_size="NOS", proposed_size=None,
    )
    assert proposed_changes(absent) == {}

    explicit_clear = SimpleNamespace(
        original_item_name="X", proposed_item_name="X", original_brand=None, proposed_brand=None,
        original_size="NOS", proposed_size="",
    )
    assert proposed_changes(explicit_clear) == {"size": {"original": "NOS", "proposed": ""}}


class TestCleanupUpload:
    def test_valid_workbook_accepted(self, client, staff_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
        path = write_fixture_workbook(tmp_path / "master.xlsx")
        with path.open("rb") as fh:
            resp = client.post(
                "/api/tools/item-combination-cleanup/runs",
                headers=staff_headers,
                files={"file": ("master.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert resp.status_code == 201, resp.text
        body = resp.json()
        assert body["status"] == "uploaded"
        assert body["row_count"] == 5
        assert body["sheet_name"] == "Item_Master_List"
        assert "ExtraNote" in body["headers"]
        assert body["validation_status"] is None

    def test_missing_column_rejected(self, client, staff_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
        from openpyxl import Workbook
        path = tmp_path / "bad.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["Item_Id", "Barcode", "Item Name", "Brand", "Size"])
        ws.append(["1", "A", "X", None, None])
        wb.save(path)
        with path.open("rb") as fh:
            resp = client.post(
                "/api/tools/item-combination-cleanup/runs",
                headers=staff_headers,
                files={"file": ("bad.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert resp.status_code == 422
        assert "HsnCode" in resp.json()["detail"]

    def test_duplicate_item_id_rejected(self, client, staff_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
        from tests.fixtures.item_combination_cleanup import mutate_column
        path = write_fixture_workbook(tmp_path / "dup.xlsx")
        mutate_column(path, "Item_Id", 3, "1001")
        with path.open("rb") as fh:
            resp = client.post(
                "/api/tools/item-combination-cleanup/runs",
                headers=staff_headers,
                files={"file": ("dup.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert resp.status_code == 422
        assert "Duplicate Item_Id" in resp.json()["detail"]

    def test_non_xlsx_rejected(self, client, staff_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
        path = tmp_path / "notes.csv"
        path.write_text("a,b\n1,2\n")
        with path.open("rb") as fh:
            resp = client.post(
                "/api/tools/item-combination-cleanup/runs",
                headers=staff_headers,
                files={"file": ("notes.csv", fh, "text/csv")},
            )
        assert resp.status_code == 400

    def test_requires_auth(self, client, tmp_path):
        path = write_fixture_workbook(tmp_path / "master.xlsx")
        with path.open("rb") as fh:
            resp = client.post(
                "/api/tools/item-combination-cleanup/runs",
                files={"file": ("master.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert resp.status_code == 401


class TestCleanupExportDownload:
    def test_upload_export_download_round_trip(self, client, staff_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
        path = write_fixture_workbook(tmp_path / "master.xlsx")

        with test_engine.connect() as conn:
            before = conn.execute(text("SELECT COUNT(*) FROM raw.raw_item_combinations")).scalar()

        with path.open("rb") as fh:
            created = client.post(
                "/api/tools/item-combination-cleanup/runs",
                headers=staff_headers,
                files={"file": ("master.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert created.status_code == 201, created.text
        run_id = created.json()["id"]

        listed = client.get("/api/tools/item-combination-cleanup/runs", headers=staff_headers)
        assert listed.status_code == 200
        assert any(r["id"] == run_id for r in listed.json())

        got = client.get(f"/api/tools/item-combination-cleanup/runs/{run_id}", headers=staff_headers)
        assert got.status_code == 200
        assert got.json()["status"] == "uploaded"

        exported = client.post(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/export",
            headers=staff_headers,
        )
        assert exported.status_code == 200, exported.text
        body = exported.json()
        assert body["status"] == "exported"
        assert body["validation_status"] == "passed"
        assert body["has_output"] is True

        download = client.get(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/download",
            headers=staff_headers,
        )
        assert download.status_code == 200
        assert download.content[:2] == b"PK"
        assert "verified.xlsx" in download.headers.get("content-disposition", "")

        with test_engine.connect() as conn:
            after = conn.execute(text("SELECT COUNT(*) FROM raw.raw_item_combinations")).scalar()
        assert after == before

    def test_download_before_export_404(self, client, staff_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
        path = write_fixture_workbook(tmp_path / "master.xlsx")
        with path.open("rb") as fh:
            created = client.post(
                "/api/tools/item-combination-cleanup/runs",
                headers=staff_headers,
                files={"file": ("master.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        run_id = created.json()["id"]
        resp = client.get(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/download",
            headers=staff_headers,
        )
        assert resp.status_code == 404


def _seed_phase2_purchases() -> None:
    with test_engine.begin() as conn:
        conn.execute(text("DELETE FROM raw.raw_purchase_itemwise WHERE barcode LIKE 'ICC2_%'"))
        batch = str(uuid.uuid4())
        for purchase in PHASE2_PURCHASES:
            conn.execute(
                text("""
                    INSERT INTO raw.raw_purchase_itemwise
                        (import_batch_id, source_file_name, barcode, mrp, supplier_name_raw,
                         purchase_date_raw, purchase_id, invoice_no)
                    VALUES
                        (:bid, 'icc_phase2_fixture.xlsx', :bc, :mrp, :sup, :dt, :pid, :inv)
                """),
                {
                    "bid": batch,
                    "bc": purchase["barcode"],
                    "mrp": purchase["mrp"],
                    "sup": purchase["supplier"],
                    "dt": purchase["purchase_date"],
                    "pid": purchase["purchase_id"],
                    "inv": purchase["invoice_no"],
                },
            )


class TestCleanupProcess:
    def test_process_proposals_and_invariants(self, client, staff_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
        _seed_phase2_purchases()
        path = write_phase2_workbook(tmp_path / "phase2.xlsx")
        original_bytes = path.read_bytes()

        with test_engine.connect() as conn:
            before_raw = conn.execute(text("SELECT COUNT(*) FROM raw.raw_item_combinations")).scalar()
            before_purchases = conn.execute(
                text("SELECT COUNT(*) FROM raw.raw_purchase_itemwise WHERE barcode LIKE 'ICC2_%'")
            ).scalar()

        with path.open("rb") as fh:
            created = client.post(
                "/api/tools/item-combination-cleanup/runs",
                headers=staff_headers,
                files={"file": ("phase2.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert created.status_code == 201, created.text
        run_id = created.json()["id"]
        stored = tmp_path / "runs" / run_id / "source.xlsx"
        stored_bytes = stored.read_bytes()

        processed = client.post(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/process",
            headers=staff_headers,
        )
        assert processed.status_code == 200, processed.text
        body = processed.json()
        assert body["status"] == "processed"
        summary = body["summary"]
        assert summary["row_count"] == 10
        assert summary["product_type"] == {
            "LOOSE": 2,
            "PACKED": 1,
            "PURCHASED": 6,
            "UNKNOWN": 1,
        }
        assert stored.read_bytes() == stored_bytes
        assert original_bytes == path.read_bytes()

        again = client.post(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/process",
            headers=staff_headers,
        )
        assert again.status_code == 200
        assert again.json()["summary"]["product_type"] == summary["product_type"]

        with test_engine.connect() as conn:
            rows = conn.execute(
                text("""
                    SELECT item_id_key, product_type, proposed_item_name, proposed_brand,
                           proposed_size, classification_confidence, review_status,
                           supplier_name, supplier_match_method, source_identity,
                           original_hsn, original_brand
                    FROM app.item_combination_cleanup_rows
                    WHERE run_id = :rid
                    ORDER BY item_id_key
                """),
                {"rid": run_id},
            ).mappings().all()
            after_raw = conn.execute(text("SELECT COUNT(*) FROM raw.raw_item_combinations")).scalar()
            after_purchases = conn.execute(
                text("SELECT COUNT(*) FROM raw.raw_purchase_itemwise WHERE barcode LIKE 'ICC2_%'")
            ).scalar()
            row_count = conn.execute(
                text("SELECT COUNT(*) FROM app.item_combination_cleanup_rows WHERE run_id = :rid"),
                {"rid": run_id},
            ).scalar()

        assert after_raw == before_raw
        assert after_purchases == before_purchases
        assert row_count == 10
        by_id = {row["item_id_key"]: row for row in rows}

        loose = by_id["P2-1001"]
        assert loose["product_type"] == "LOOSE"
        assert loose["proposed_item_name"] == "SAME LOOSE"
        assert loose["proposed_brand"] == "SRI BALAJI TRADERS"
        assert loose["classification_confidence"] == "HIGH"
        assert loose["supplier_match_method"] == "barcode_mrp"

        packed = by_id["P2-1002"]
        assert packed["product_type"] == "PACKED"
        assert packed["proposed_item_name"] == "SAME 1KG"
        assert packed["proposed_brand"] == "PKG"
        assert packed["proposed_size"] == "1KG"
        assert packed["classification_confidence"] == "HIGH"
        assert packed["review_status"] == "NEEDS_REVIEW"

        unknown = by_id["P2-1003"]
        assert unknown["product_type"] == "UNKNOWN"
        assert unknown["proposed_item_name"] == "GREEN TEA 100G"
        assert unknown["proposed_size"] == "100G"
        assert unknown["proposed_brand"] is None

        pkg_buy = by_id["P2-1004"]
        assert pkg_buy["product_type"] == "PURCHASED"
        assert pkg_buy["proposed_brand"] == "OIL DISTRIBUTORS"
        assert pkg_buy["original_brand"] == "PKG"

        phm = by_id["P2-1005"]
        assert phm["product_type"] == "PURCHASED"
        assert phm["proposed_brand"] == "EXTERNAL MILLERS"
        assert phm["proposed_brand"] != "PHM"

        adu = by_id["P2-1006"]
        assert adu["product_type"] == "PURCHASED"
        assert adu["proposed_brand"] == "ADUKALE FOODS"

        loose_none = by_id["P2-1007"]
        assert loose_none["product_type"] == "LOOSE"
        assert loose_none["proposed_brand"] == "KEEPME"
        assert loose_none["review_status"] == "NEEDS_REVIEW"

        shared_a = by_id["P2-1008"]
        shared_b = by_id["P2-1009"]
        assert shared_a["product_type"] == "PURCHASED"
        assert shared_b["product_type"] == "PURCHASED"
        assert shared_a["classification_confidence"] == "LOW"
        assert shared_b["classification_confidence"] == "LOW"

        fallback = by_id["P2-1010"]
        assert fallback["product_type"] == "PURCHASED"
        assert fallback["supplier_match_method"] == "barcode"
        assert fallback["classification_confidence"] == "HIGH"
        assert fallback["review_status"] == "AUTO_APPROVED"

        for row in rows:
            identity = row["source_identity"]
            assert set(identity) == {
                "item_id", "barcode", "mrp", "rate", "purchase_price", "expiry", "stock",
            }
            assert "proposed_item_id" not in identity
            assert row["original_hsn"] is not None


class TestCleanupInspection:
    def _processed_run(self, client, staff_headers, tmp_path, monkeypatch):
        monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
        _seed_phase2_purchases()
        path = write_phase2_workbook(tmp_path / "phase2.xlsx")
        with path.open("rb") as fh:
            created = client.post(
                "/api/tools/item-combination-cleanup/runs",
                headers=staff_headers,
                files={"file": ("phase2.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        run_id = created.json()["id"]
        processed = client.post(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/process",
            headers=staff_headers,
        )
        assert processed.status_code == 200, processed.text
        return run_id, tmp_path / "runs" / run_id / "source.xlsx"

    def test_list_filter_detail_and_audit_export(self, client, staff_headers, tmp_path, monkeypatch):
        run_id, source = self._processed_run(client, staff_headers, tmp_path, monkeypatch)
        before = source.read_bytes()

        listed = client.get(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/rows",
            headers=staff_headers,
        )
        assert listed.status_code == 200, listed.text
        body = listed.json()
        assert body["total"] == 10
        assert body["stats"]["product_type"]["PACKED"] == 1
        assert body["stats"]["product_type"]["LOOSE"] == 2
        assert body["stats"]["brand_changed"] >= 1
        assert {row["item_id"] for row in body["items"]} == {
            "P2-1001", "P2-1002", "P2-1003", "P2-1004", "P2-1005",
            "P2-1006", "P2-1007", "P2-1008", "P2-1009", "P2-1010",
        }

        packed = client.get(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/rows",
            headers=staff_headers,
            params={"type": "PACKED"},
        )
        assert packed.json()["total"] == 1
        assert packed.json()["items"][0]["proposed_brand"] == "PKG"

        brand_changed = client.get(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/rows",
            headers=staff_headers,
            params={"brand_changed": True, "type": "PURCHASED"},
        )
        assert brand_changed.status_code == 200
        assert brand_changed.json()["total"] >= 1

        search = client.get(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/rows",
            headers=staff_headers,
            params={"q": "SAME LOOSE"},
        )
        assert search.json()["total"] == 1

        detail = client.get(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/rows/P2-1001",
            headers=staff_headers,
        )
        assert detail.status_code == 200
        evidence = detail.json()
        assert evidence["product_type"] == "LOOSE"
        assert evidence["supplier_name"] == "SRI BALAJI TRADERS"
        assert evidence["classification_evidence"]["match_method"] == "barcode_mrp"
        assert "rules" in evidence["classification_evidence"]

        missing = client.get(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/rows/NOPE",
            headers=staff_headers,
        )
        assert missing.status_code == 404

        exported = client.post(
            f"/api/tools/item-combination-cleanup/runs/{run_id}/proposals-export",
            headers=staff_headers,
        )
        assert exported.status_code == 200
        assert exported.content[:2] == b"PK"
        assert "proposals_audit.xlsx" in exported.headers.get("content-disposition", "")
        assert source.read_bytes() == before

        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(exported.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert headers[:5] == [
            "Item_Id", "Barcode", "MRP", "Original Item Name", "Proposed Item Name",
        ]
        assert "Classification Evidence" in headers
        assert ws.max_row == 11
        wb.close()

        with test_engine.connect() as conn:
            raw_count = conn.execute(text("SELECT COUNT(*) FROM raw.raw_item_combinations")).scalar()
        assert raw_count is not None


class TestCleanupApproval:
    def test_bulk_decision_approves_exact_selected_rows(self, client, staff_headers, tmp_path, monkeypatch):
        run_id, _ = TestCleanupInspection()._processed_run(client, staff_headers, tmp_path, monkeypatch)
        base = f"/api/tools/item-combination-cleanup/runs/{run_id}"
        response = client.post(f"{base}/bulk-decision", headers=staff_headers, json={"status": "APPROVED", "item_ids": ["P2-1002", "P2-1003"]})
        assert response.status_code == 200, response.text
        assert response.json()["updated"] == 2
    def test_bulk_approval_requires_filter_and_rejects_low_confidence(self, client, staff_headers, tmp_path, monkeypatch):
        run_id, _ = TestCleanupInspection()._processed_run(client, staff_headers, tmp_path, monkeypatch)
        base = f"/api/tools/item-combination-cleanup/runs/{run_id}"
        assert client.post(f"{base}/bulk-approval", headers=staff_headers, json={}).status_code == 422
        low = client.post(f"{base}/bulk-approval", headers=staff_headers, json={"confidence": "LOW"})
        assert low.status_code == 422
        approved = client.post(f"{base}/bulk-approval", headers=staff_headers, json={"type": "PACKED"})
        assert approved.status_code == 200, approved.text
        assert approved.json()["updated"] == 1
        # Reprocessing would delete the stored human decision, so it is blocked.
        assert client.post(f"{base}/process", headers=staff_headers).status_code == 409

    def test_approved_export_applies_explicit_packed_but_not_pending_brand_change(self, client, staff_headers, tmp_path, monkeypatch):
        run_id, source = TestCleanupInspection()._processed_run(client, staff_headers, tmp_path, monkeypatch)
        base = f"/api/tools/item-combination-cleanup/runs/{run_id}"
        approved = client.post(f"{base}/rows/P2-1002/approval", headers=staff_headers, json={"status": "APPROVED"})
        assert approved.status_code == 200, approved.text
        exported = client.post(f"{base}/approved-export", headers=staff_headers)
        assert exported.status_code == 200, exported.text
        download = client.get(f"{base}/approved-download", headers=staff_headers)
        assert download.status_code == 200
        from io import BytesIO
        from openpyxl import load_workbook
        src, out = load_workbook(source), load_workbook(BytesIO(download.content))
        headers = {cell.value: cell.column for cell in src.active[1]}
        # P2-1002 is explicitly approved: Brand PKG is applied.
        assert out.active.cell(3, headers["Brand"]).value == "PKG"
        # P2-1001 has an unapproved supplier-brand proposal: source Brand remains unchanged.
        assert out.active.cell(2, headers["Brand"]).value == src.active.cell(2, headers["Brand"]).value
        src.close(); out.close()


class TestStaffReviewImport:
    @staticmethod
    def _workbook(run_id, rows):
        from io import BytesIO
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.title = "Review Items"
        ws.append(["Review ID", "Item_Id", "Staff Decision", "Correct Item Name", "Correct Brand", "Correct Size"])
        for item_id, decision, name, brand, size in rows:
            ws.append([f"{run_id}:{item_id}", item_id, decision, name, brand, size])
        out = BytesIO(); wb.save(out); return out.getvalue()

    def test_preview_then_apply_protects_resolved_rows(self, client, staff_headers, tmp_path, monkeypatch):
        run_id, _ = TestCleanupInspection()._processed_run(client, staff_headers, tmp_path, monkeypatch)
        base = f"/api/tools/item-combination-cleanup/runs/{run_id}"
        with test_engine.begin() as conn:
            conn.execute(text("UPDATE app.item_combination_cleanup_rows SET approval_status='APPROVED' WHERE run_id=:rid AND item_id_key='P2-1002'"), {"rid": run_id})
        content = self._workbook(run_id, [("P2-1002", "NO", None, None, None), ("P2-1003", "YES", None, None, None)])
        files = {"file": ("staff.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        preview = client.post(f"{base}/staff-review-import-preview", headers=staff_headers, files=files)
        assert preview.status_code == 200, preview.text
        assert preview.json()["summary"] == {"applicable": 1, "protected": 1, "not_sure": 0, "invalid": 0}
        applied = client.post(f"{base}/staff-review-import-apply", headers=staff_headers, data={"preview_hash": preview.json()["file_hash"]}, files=files)
        assert applied.status_code == 200, applied.text
        assert applied.json()["applied_rows"] == 1
        assert applied.json()["protected_rows"] == 1
        with test_engine.connect() as conn:
            protected = conn.execute(text("SELECT approval_status FROM app.item_combination_cleanup_rows WHERE run_id=:rid AND item_id_key='P2-1002'"), {"rid": run_id}).scalar()
            imported = conn.execute(text("SELECT count(*) FROM app.item_combination_cleanup_field_decisions WHERE run_id=:rid AND item_id_key='P2-1003' AND review_source='STAFF_IMPORT'"), {"rid": run_id}).scalar()
        assert protected == "APPROVED"
        assert imported == 2
