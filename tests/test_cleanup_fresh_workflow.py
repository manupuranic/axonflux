from __future__ import annotations

import uuid
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from api.tools.item_combination_cleanup.models import (
    ItemCombinationCleanupFieldDecision,
    ItemCombinationCleanupNameSuggestion,
    ItemCombinationCleanupRow,
    ItemCombinationCleanupRun,
)
from api.tools.item_combination_cleanup.name_semantic_eval import NameChange, NameSemanticOutput
from api.tools.item_combination_cleanup.semantic import SemanticAssessmentOutput
from tests.conftest import TestSessionLocal
from tests.fixtures.item_combination_cleanup import HEADERS, write_phase2_workbook


def _write_fresh_fixture(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Item_Master_List"
    sheet.append(HEADERS)
    sheet.append(["REUSE-EDIT", 1, "BC-EDIT", "HEAD SHOULDER NEEM SHMOP 650ML", "33051090", "GST(18)", None, None])
    sheet.append(["STALE-EDIT", 1, "BC-STALE", "HEAD SHOULDER NEEM SHAMPOO 650ML NEW", "33051090", "GST(18)", None, None])
    sheet.append(["REUSE-REJECT", 1, "BC-REJECT", "COMFORT ROYALE 4/-", "34025000", "GST(18)", None, None])
    sheet.append(["REUSE-SIZE", 1, "BC-SIZE", "GREEN TEA 100GM", "09021000", "GST(5)", None, None])
    sheet.append(["STALE-SIZE", 1, "BC-STALE-SIZE", "GREEN TEA 100GM", "09021000", "GST(5)", None, "200G"])
    sheet.append(["NEW-SHMP", 1, "BC-NEW", "GARNIER SHMP 100ML", "33051090", "GST(18)", None, None])
    for index in range(12):
        sheet.append([f"COMMON-{index}", 1, f"BC-COMMON-{index}", f"FRESH BANANA LIQUID BUNCH {index}", "08039010", "GST(5)", None, None])
    sheet.append(["SEMANTIC-NEW", 1, "BC-SEMANTIC", "FRESH BANAN BUNCH", "08039010", "GST(5)", None, None])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _seed_historical_run() -> uuid.UUID:
    with TestSessionLocal() as db:
        run = ItemCombinationCleanupRun(
            id=uuid.uuid4(),
            source_file_name="historical.xlsx",
            original_path="historical.xlsx",
            status="processed",
        )
        db.add(run)
        db.flush()
        rows = [
            ItemCombinationCleanupRow(
                run_id=run.id,
                row_index=index,
                item_id_key=item_id,
                source_identity={"item_id": item_id, "barcode": barcode},
                original_item_name=name,
                original_size=original_size,
                proposed_item_name=name,
                proposed_size=proposed_size,
                product_type="UNKNOWN",
                review_status="NEEDS_REVIEW",
                approval_status="APPROVED",
            )
            for index, item_id, barcode, name, original_size, proposed_size in (
                (2, "REUSE-EDIT", "BC-EDIT", "HEAD SHOULDER NEEM SHMOP 650ML", None, "650ML"),
                (3, "STALE-EDIT", "BC-STALE", "HEAD SHOULDER NEEM SHMOP 650ML", None, "650ML"),
                (4, "REUSE-REJECT", "BC-REJECT", "COMFORT ROYALE 4/-", None, None),
                (5, "REUSE-SIZE", "BC-SIZE", "GREEN TEA 100GM", None, "100G"),
                (6, "STALE-SIZE", "BC-STALE-SIZE", "GREEN TEA 100GM", None, "100G"),
            )
        ]
        db.add_all(rows)
        db.flush()
        db.add_all(
            [
                ItemCombinationCleanupFieldDecision(
                    run_id=run.id,
                    item_id_key="REUSE-EDIT",
                    field_name="name",
                    original_value="HEAD SHOULDER NEEM SHMOP 650ML",
                    deterministic_proposed_value="HEAD SHOULDER NEEM SHMP 650ML",
                    edited_value="HEAD & SHOULDER NEEM SHAMPOO 650ML",
                    decision="ACCEPTED",
                    review_source="NAME_SUGGESTION:SEMANTIC_V2",
                ),
                ItemCombinationCleanupFieldDecision(
                    run_id=run.id,
                    item_id_key="STALE-EDIT",
                    field_name="name",
                    original_value="HEAD SHOULDER NEEM SHMOP 650ML",
                    deterministic_proposed_value="HEAD SHOULDER NEEM SHMP 650ML",
                    edited_value="HEAD & SHOULDER NEEM SHAMPOO 650ML",
                    decision="ACCEPTED",
                    review_source="NAME_SUGGESTION:SEMANTIC_V2",
                ),
                ItemCombinationCleanupFieldDecision(
                    run_id=run.id,
                    item_id_key="REUSE-SIZE",
                    field_name="size",
                    original_value=None,
                    deterministic_proposed_value="100G",
                    decision="ACCEPTED",
                    review_source="UI",
                ),
                ItemCombinationCleanupFieldDecision(
                    run_id=run.id,
                    item_id_key="STALE-SIZE",
                    field_name="size",
                    original_value=None,
                    deterministic_proposed_value="100G",
                    decision="ACCEPTED",
                    review_source="UI",
                ),
            ]
        )
        db.add_all(
            [
                ItemCombinationCleanupNameSuggestion(
                    run_id=run.id,
                    item_id_key="REUSE-EDIT",
                    suggestion_source="SEMANTIC_V2",
                    suggestion_version="phase5b-v2-advisory",
                    category="SEMANTIC_ADVISORY",
                    base_value="HEAD SHOULDER NEEM SHMOP 650ML",
                    suggested_value="HEAD SHOULDER NEEM SHMP 650ML",
                    transformations=[{"change_type": "SPELLING", "original_text": "SHMOP", "replacement_text": "SHMP"}],
                    evidence={"confidence": "HIGH"},
                    status="ACCEPTED",
                ),
                ItemCombinationCleanupNameSuggestion(
                    run_id=run.id,
                    item_id_key="REUSE-REJECT",
                    suggestion_source="SEMANTIC_V2",
                    suggestion_version="phase5b-v2-advisory",
                    category="SEMANTIC_ADVISORY",
                    base_value="COMFORT ROYALE 4/-",
                    suggested_value="COMFORT ROYAL 4/-",
                    transformations=[{"change_type": "SPELLING", "original_text": "ROYALE", "replacement_text": "ROYAL"}],
                    evidence={"confidence": "HIGH"},
                    status="REJECTED",
                ),
            ]
        )
        db.commit()
        return run.id


def test_fresh_upload_analysis_reuses_only_exact_historical_baselines(
    client, staff_headers, tmp_path, monkeypatch, seed_test_db
):
    monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
    historical_run_id = _seed_historical_run()
    fixture = _write_fresh_fixture(tmp_path / "fresh.xlsx")

    with fixture.open("rb") as handle:
        created = client.post(
            "/api/tools/item-combination-cleanup/runs",
            headers=staff_headers,
            files={"file": ("fresh.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert created.status_code == 201, created.text
    fresh_run_id = created.json()["id"]
    assert fresh_run_id != str(historical_run_id)

    analyzed = client.post(
        f"/api/tools/item-combination-cleanup/runs/{fresh_run_id}/analyze",
        headers=staff_headers,
        json={"include_semantic": False},
    )
    assert analyzed.status_code == 200, analyzed.text
    assert analyzed.json()["status"] == "ready_for_review"
    analysis = analyzed.json()["summary"]["analysis"]
    assert analysis["historical_field_decisions_reused"] == 1
    assert analysis["historical_name_acceptances_reused"] == 1
    assert analysis["historical_name_rejections_reused"] == 1
    assert analysis["historical_decisions_invalidated"] >= 1
    assert analysis["historical_items_matched"] == 5
    assert analysis["new_items"] == 14
    assert analysis["semantic_calls"] == 0

    name_review = client.get(
        f"/api/tools/item-combination-cleanup/runs/{fresh_run_id}/name-review",
        headers=staff_headers,
    )
    assert name_review.status_code == 200
    pending_ids = {item["item_id"] for item in name_review.json()["items"]}
    assert "NEW-SHMP" in pending_ids
    assert "REUSE-EDIT" not in pending_ids
    assert "REUSE-REJECT" not in pending_ids
    assert "STALE-EDIT" not in pending_ids

    deterministic = next(
        item for item in name_review.json()["items"] if item["item_id"] == "NEW-SHMP"
    )["suggestions"][0]
    accepted = client.post(
        f"/api/tools/item-combination-cleanup/runs/{fresh_run_id}/name-suggestions/{deterministic['id']}/decision",
        headers=staff_headers,
        json={"decision": "ACCEPTED"},
    )
    assert accepted.status_code == 200, accepted.text
    exported = client.post(
        f"/api/tools/item-combination-cleanup/runs/{fresh_run_id}/approved-export",
        headers=staff_headers,
    )
    assert exported.status_code == 200, exported.text
    downloaded = client.get(
        f"/api/tools/item-combination-cleanup/runs/{fresh_run_id}/approved-download",
        headers=staff_headers,
    )
    workbook = load_workbook(BytesIO(downloaded.content), data_only=False)
    sheet = workbook["Item_Master_List"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    exported_names = {
        str(sheet.cell(row, headers["Item_Id"]).value): sheet.cell(row, headers["Item Name"]).value
        for row in range(2, sheet.max_row + 1)
    }
    assert exported_names["NEW-SHMP"] == "GARNIER SHAMPOO 100ML"
    workbook.close()

    with TestSessionLocal() as db:
        historical_decisions = db.query(ItemCombinationCleanupFieldDecision).filter_by(run_id=historical_run_id).count()
        historical_suggestions = db.query(ItemCombinationCleanupNameSuggestion).filter_by(run_id=historical_run_id).count()
        assert historical_decisions == 4
        assert historical_suggestions == 2


def test_fresh_analysis_runs_existing_semantic_v2_advisory_through_api(
    client, staff_headers, tmp_path, monkeypatch, seed_test_db
):
    from api.tools.item_combination_cleanup import workflow

    monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
    _seed_historical_run()
    fixture = _write_fresh_fixture(tmp_path / "semantic-fresh.xlsx")
    calls = []

    def fake_evaluate(payload, **_kwargs):
        calls.append(payload)
        return (
            NameSemanticOutput(
                result="CORRECTION",
                confidence="HIGH",
                reason="Catalog siblings consistently use BANANA.",
                suggested_name="FRESH BANANA BUNCH",
                changes=[
                    NameChange(
                        original_text="BANAN",
                        replacement_text="BANANA",
                        change_type="SPELLING",
                    )
                ],
            ),
            SimpleNamespace(prompt_tokens=10, completion_tokens=5),
        )

    monkeypatch.setattr(workflow, "evaluate_name", fake_evaluate, raising=False)
    with fixture.open("rb") as handle:
        created = client.post(
            "/api/tools/item-combination-cleanup/runs",
            headers=staff_headers,
            files={"file": ("semantic-fresh.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    run_id = created.json()["id"]
    analyzed = client.post(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/analyze",
        headers=staff_headers,
        json={"include_semantic": True},
    )

    assert analyzed.status_code == 200, analyzed.text
    analysis = analyzed.json()["summary"]["analysis"]
    assert analysis["semantic_calls"] == 1
    assert analysis["semantic_suggestions_created"] == 1
    assert analysis["historical_name_rejections_reused"] == 1
    assert [call["candidate"]["item_id"] for call in calls] == ["SEMANTIC-NEW"]
    review = client.get(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/name-review",
        headers=staff_headers,
    ).json()
    semantic = next(item for item in review["items"] if item["item_id"] == "SEMANTIC-NEW")
    assert semantic["suggestions"][0]["source"] == "SEMANTIC_V2"
    assert semantic["suggestions"][0]["suggested_value"] == "FRESH BANANA BUNCH"


def test_fresh_analysis_runs_existing_packed_semantic_advisory_through_api(
    client, staff_headers, tmp_path, monkeypatch, seed_test_db
):
    from api.tools.item_combination_cleanup import workflow
    from tests.test_api_item_combination_cleanup import _seed_phase2_purchases

    monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))
    _seed_phase2_purchases()
    fixture = write_phase2_workbook(tmp_path / "packed-fresh.xlsx")
    calls = []

    def fake_packed_evaluate(payload, **_kwargs):
        calls.append(payload)
        return (
            SemanticAssessmentOutput(
                relationship="LIKELY_PACKED_VERSION",
                confidence="HIGH",
                reason="The bounded evidence supports a packed relationship.",
                signals_for_packed=["Matched loose source"],
                signals_against_packed=[],
                identity_interpretations=[],
            ),
            SimpleNamespace(prompt_tokens=12, completion_tokens=6),
        )

    monkeypatch.setattr(workflow, "evaluate_packed", fake_packed_evaluate, raising=False)
    with fixture.open("rb") as handle:
        created = client.post(
            "/api/tools/item-combination-cleanup/runs",
            headers=staff_headers,
            files={"file": ("packed-fresh.xlsx", handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    run_id = created.json()["id"]
    analyzed = client.post(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/analyze",
        headers=staff_headers,
        json={"include_semantic": True},
    )

    assert analyzed.status_code == 200, analyzed.text
    analysis = analyzed.json()["summary"]["analysis"]
    assert analysis["packed_semantic_calls"] == 1
    assert analysis["packed_semantic_assessments_created"] == 1
    assert len(calls) == 1
    packed = client.get(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/rows/P2-1002",
        headers=staff_headers,
    ).json()
    assert packed["semantic_assessment"]["relationship"] == "LIKELY_PACKED_VERSION"
    assert packed["proposed_brand"] == "PKG"
