from __future__ import annotations

import uuid
from datetime import datetime, timedelta, timezone

from api.tools.item_combination_cleanup.models import (
    ItemCombinationCleanupFieldDecision,
    ItemCombinationCleanupNameSuggestion,
    ItemCombinationCleanupRow,
    ItemCombinationCleanupRun,
)
from api.tools.item_combination_cleanup.service import effective_name_for_phase5
from tests.conftest import TestSessionLocal
from tests.fixtures.item_combination_cleanup import HEADERS


def _run(db) -> ItemCombinationCleanupRun:
    run = ItemCombinationCleanupRun(
        id=uuid.uuid4(),
        source_file_name="name-review.xlsx",
        original_path="name-review.xlsx",
        status="processed",
    )
    db.add(run)
    db.flush()
    return run


def _row(db, run, item_id: str, name: str = "DARK CHOCOLT BAR"):
    row = ItemCombinationCleanupRow(
        run_id=run.id,
        row_index=int(item_id.split("-")[-1]) if item_id.split("-")[-1].isdigit() else 1,
        item_id_key=item_id,
        source_identity={"item_id": item_id, "barcode": f"BC-{item_id}"},
        original_item_name=name,
        original_brand="Protected Brand",
        original_size="Protected Size",
        proposed_brand="Protected Brand Override",
        proposed_size="Protected Size Override",
        product_type="PACKED",
        supplier_name="Protected Supplier",
        review_status="NEEDS_REVIEW",
        approval_status="PENDING",
    )
    db.add(row)
    return row


def _suggestion(
    db,
    run,
    item_id: str,
    *,
    source: str,
    base: str,
    proposed: str,
    transformation,
    status: str = "PENDING",
):
    suggestion = ItemCombinationCleanupNameSuggestion(
        run_id=run.id,
        item_id_key=item_id,
        field_name="name",
        suggestion_source=source,
        suggestion_version=(
            "phase5b-v2-effective-baseline" if source == "NAMING_STANDARD" else "phase5b-v2-advisory"
        ),
        category="HUMAN_VERIFIED_CORRECTION" if source == "NAMING_STANDARD" else "SEMANTIC_ADVISORY",
        base_value=base,
        suggested_value=proposed,
        transformations=[transformation],
        evidence=(
            {"categories": ["HUMAN_VERIFIED_CORRECTION"]}
            if source == "NAMING_STANDARD"
            else {
                "confidence": "HIGH",
                "reason": "Same-item evidence supports this correction.",
                "payload": {
                    "candidate": {
                        "detector_reason": "RARE_NEAR_NEIGHBOR",
                        "corroborating_evidence": "PURCHASE",
                    },
                    "context": {
                        "purchase_item_names": [proposed],
                        "catalog_siblings": [],
                    },
                },
            }
        ),
        status=status,
    )
    db.add(suggestion)
    return suggestion


def test_name_review_list_filters_stale_and_decided_and_classifies_overlap(
    client, staff_headers, seed_test_db
):
    with TestSessionLocal() as db:
        run = _run(db)
        _row(db, run, "det-1")
        _suggestion(
            db, run, "det-1", source="NAMING_STANDARD", base="DARK CHOCOLT BAR",
            proposed="DARK CHOCOLATE BAR", transformation="CHOCOLT → CHOCOLATE",
        )
        _row(db, run, "sem-1", "CADBERY COOKIE")
        _suggestion(
            db, run, "sem-1", source="SEMANTIC_V2", base="CADBERY COOKIE",
            proposed="CADBURY COOKIE",
            transformation={"original_text": "CADBERY", "replacement_text": "CADBURY", "change_type": "SPELLING"},
        )
        _row(db, run, "same-1")
        _suggestion(
            db, run, "same-1", source="NAMING_STANDARD", base="DARK CHOCOLT BAR",
            proposed="DARK CHOCOLATE BAR", transformation="CHOCOLT → CHOCOLATE",
        )
        _suggestion(
            db, run, "same-1", source="SEMANTIC_V2", base="DARK CHOCOLT BAR",
            proposed="DARK CHOCOLATE BAR",
            transformation={"original_text": "CHOCOLT", "replacement_text": "CHOCOLATE", "change_type": "SPELLING"},
        )
        _row(db, run, "conflict-1")
        _suggestion(
            db, run, "conflict-1", source="NAMING_STANDARD", base="DARK CHOCOLT BAR",
            proposed="DARK CHOCOLATE BAR", transformation="CHOCOLT → CHOCOLATE",
        )
        _suggestion(
            db, run, "conflict-1", source="SEMANTIC_V2", base="DARK CHOCOLT BAR",
            proposed="DARK COCOA BAR",
            transformation={"original_text": "CHOCOLT", "replacement_text": "COCOA", "change_type": "OTHER"},
        )
        _row(db, run, "stale-status")
        _suggestion(
            db, run, "stale-status", source="NAMING_STANDARD", base="DARK CHOCOLT BAR",
            proposed="DARK CHOCOLATE BAR", transformation="CHOCOLT → CHOCOLATE",
            status="SUPERSEDED_STALE_BASELINE",
        )
        _row(db, run, "stale-base", "CADBERY COOKIE 100G")
        stale_baseline = _suggestion(
            db, run, "stale-base", source="SEMANTIC_V2", base="CADBERY COOKIE 100GM",
            proposed="CADBURY COOKIE 100GM",
            transformation={"original_text": "CADBERY", "replacement_text": "CADBURY", "change_type": "SPELLING"},
        )
        decided = _row(db, run, "decided-1")
        decided_suggestion = _suggestion(
            db, run, "decided-1", source="NAMING_STANDARD", base="DARK CHOCOLT BAR",
            proposed="DARK CHOCOLATE BAR", transformation="CHOCOLT → CHOCOLATE",
        )
        db.flush()
        decided_suggestion.created_at = datetime.now(timezone.utc) - timedelta(days=1)
        db.add(
            ItemCombinationCleanupFieldDecision(
                run_id=run.id,
                item_id_key=decided.item_id_key,
                field_name="name",
                original_value=decided.original_item_name,
                deterministic_proposed_value="DARK CHOCOLATE BAR",
                decision="ACCEPTED",
                reviewed_at=datetime.now(timezone.utc),
            )
        )
        db.commit()
        run_id = str(run.id)
        stale_baseline_id = str(stale_baseline.id)

    response = client.get(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/name-review",
        headers=staff_headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["stats"] == {
        "actionable_suggestion_rows": 6,
        "unique_items": 4,
        "deterministic_only_items": 1,
        "semantic_only_items": 1,
        "overlap_items": 2,
        "already_decided_excluded": 1,
        "stale_excluded": 2,
    }
    items = {item["item_id"]: item for item in body["items"]}
    assert set(items) == {"det-1", "sem-1", "same-1", "conflict-1"}
    assert items["same-1"]["overlap_classification"] == "EXACT_SAME_RESULT"
    assert items["conflict-1"]["overlap_classification"] == "CONFLICTING_RESULTS"
    assert items["sem-1"]["suggestions"][0]["confidence"] == "HIGH"
    assert items["sem-1"]["suggestions"][0]["purchase_item_names"] == ["CADBURY COOKIE"]
    refused = client.post(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/name-suggestions/{stale_baseline_id}/decision",
        headers=staff_headers,
        json={"decision": "ACCEPTED"},
    )
    assert refused.status_code == 409


def _single_review_run(db, *, name="DARK CHOCOLT BAR", transformation="CHOCOLT → CHOCOLATE"):
    run = _run(db)
    row = _row(db, run, "review-1", name)
    suggestion = _suggestion(
        db, run, "review-1", source="NAMING_STANDARD", base=name,
        proposed=name.replace("CHOCOLT", "CHOCOLATE"), transformation=transformation,
    )
    db.commit()
    return run.id, suggestion.id, row.item_id_key


def test_name_suggestion_accept_edit_reject_pending_and_protected_field_isolation(
    client, staff_headers, seed_test_db
):
    # Pending/read-only listing changes nothing.
    with TestSessionLocal() as db:
        pending_run, pending_id, _ = _single_review_run(db)
    listed = client.get(
        f"/api/tools/item-combination-cleanup/runs/{pending_run}/name-review",
        headers=staff_headers,
    )
    assert listed.status_code == 200
    with TestSessionLocal() as db:
        pending = db.get(ItemCombinationCleanupNameSuggestion, pending_id)
        assert pending.status == "PENDING"
        assert db.query(ItemCombinationCleanupFieldDecision).filter_by(run_id=pending_run).count() == 0

    # Accept makes exactly the suggestion authoritative.
    with TestSessionLocal() as db:
        accept_run, accept_id, accept_item = _single_review_run(db)
    accepted = client.post(
        f"/api/tools/item-combination-cleanup/runs/{accept_run}/name-suggestions/{accept_id}/decision",
        headers=staff_headers,
        json={"decision": "ACCEPTED"},
    )
    assert accepted.status_code == 200, accepted.text
    with TestSessionLocal() as db:
        row = db.query(ItemCombinationCleanupRow).filter_by(run_id=accept_run, item_id_key=accept_item).one()
        decision = db.query(ItemCombinationCleanupFieldDecision).filter_by(run_id=accept_run, item_id_key=accept_item, field_name="name").one()
        assert effective_name_for_phase5(db, row) == "DARK CHOCOLATE BAR"
        assert decision.decision == "ACCEPTED"
        assert decision.deterministic_proposed_value == "DARK CHOCOLATE BAR"
        assert (row.original_brand, row.proposed_brand, row.original_size, row.proposed_size) == (
            "Protected Brand", "Protected Brand Override", "Protected Size", "Protected Size Override",
        )
        assert row.supplier_name == "Protected Supplier"
        assert row.product_type == "PACKED"

    # Edit accepts the operator's value as authoritative over the suggestion.
    with TestSessionLocal() as db:
        edit_run, edit_id, edit_item = _single_review_run(db)
    edited = client.post(
        f"/api/tools/item-combination-cleanup/runs/{edit_run}/name-suggestions/{edit_id}/decision",
        headers=staff_headers,
        json={"decision": "EDITED", "edited_value": "DARK CHOCOLATE BAR 100G"},
    )
    assert edited.status_code == 200, edited.text
    with TestSessionLocal() as db:
        row = db.query(ItemCombinationCleanupRow).filter_by(run_id=edit_run, item_id_key=edit_item).one()
        assert effective_name_for_phase5(db, row) == "DARK CHOCOLATE BAR 100G"

    # Reject records only suggestion review; it creates no export-authoritative decision.
    with TestSessionLocal() as db:
        reject_run, reject_id, reject_item = _single_review_run(db)
    rejected = client.post(
        f"/api/tools/item-combination-cleanup/runs/{reject_run}/name-suggestions/{reject_id}/decision",
        headers=staff_headers,
        json={"decision": "REJECTED"},
    )
    assert rejected.status_code == 200, rejected.text
    with TestSessionLocal() as db:
        row = db.query(ItemCombinationCleanupRow).filter_by(run_id=reject_run, item_id_key=reject_item).one()
        assert effective_name_for_phase5(db, row) == "DARK CHOCOLT BAR"
        assert db.query(ItemCombinationCleanupFieldDecision).filter_by(run_id=reject_run).count() == 0
        assert db.get(ItemCombinationCleanupNameSuggestion, reject_id).status == "REJECTED"


def test_bulk_name_acceptance_is_atomic_and_limited_to_safe_exact_rules(
    client, staff_headers, seed_test_db
):
    with TestSessionLocal() as db:
        run = _run(db)
        _row(db, run, "safe-1")
        safe = _suggestion(
            db, run, "safe-1", source="NAMING_STANDARD", base="DARK CHOCOLT BAR",
            proposed="DARK CHOCOLATE BAR", transformation="CHOCOLT → CHOCOLATE",
        )
        _row(db, run, "unsafe-1", "HERBAL SHMP")
        unsafe = _suggestion(
            db, run, "unsafe-1", source="NAMING_STANDARD", base="HERBAL SHMP",
            proposed="HERBAL SHAMPOO", transformation="SHMP → SHAMPOO",
        )
        db.commit()
        run_id = run.id
        safe_id, unsafe_id = safe.id, unsafe.id

    refused = client.post(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/name-suggestions/bulk-decision",
        headers=staff_headers,
        json={"suggestion_ids": [str(safe_id), str(unsafe_id)], "decision": "ACCEPTED"},
    )
    assert refused.status_code == 422, refused.text
    with TestSessionLocal() as db:
        assert db.get(ItemCombinationCleanupNameSuggestion, safe_id).status == "PENDING"
        assert db.query(ItemCombinationCleanupFieldDecision).filter_by(run_id=run_id).count() == 0

    accepted = client.post(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/name-suggestions/bulk-decision",
        headers=staff_headers,
        json={"suggestion_ids": [str(safe_id)], "decision": "ACCEPTED"},
    )
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["updated"] == 1


def test_approved_export_applies_only_the_accepted_name_suggestion(
    client, staff_headers, seed_test_db, tmp_path, monkeypatch
):
    from io import BytesIO
    from openpyxl import Workbook, load_workbook

    source = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Item_Master_List"
    ws.append(HEADERS)
    original = [
        "1001", 1, "8900000000001", "DARK CHOCOLT BAR", "33051090", "GST(18)",
        "Protected Brand", "Protected Size", None, None, None, 1, 9, 40, 55, 50, 10,
        "protected-extra",
    ]
    ws.append(original)
    wb.save(source)
    wb.close()
    monkeypatch.setenv("ITEM_COMBINATION_CLEANUP_DIR", str(tmp_path / "runs"))

    with TestSessionLocal() as db:
        run = _run(db)
        run.original_path = str(source)
        row = _row(db, run, "1001")
        row.source_identity = {"item_id": "1001", "barcode": "8900000000001"}
        suggestion = _suggestion(
            db, run, "1001", source="NAMING_STANDARD", base="DARK CHOCOLT BAR",
            proposed="DARK CHOCOLATE BAR", transformation="CHOCOLT → CHOCOLATE",
        )
        db.commit()
        run_id, suggestion_id = run.id, suggestion.id

    accepted = client.post(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/name-suggestions/{suggestion_id}/decision",
        headers=staff_headers,
        json={"decision": "ACCEPTED"},
    )
    assert accepted.status_code == 200, accepted.text
    exported = client.post(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/approved-export",
        headers=staff_headers,
    )
    assert exported.status_code == 200, exported.text
    downloaded = client.get(
        f"/api/tools/item-combination-cleanup/runs/{run_id}/approved-download",
        headers=staff_headers,
    )
    out = load_workbook(BytesIO(downloaded.content))
    values = [cell.value for cell in out.active[2]]
    out.close()
    assert values[3] == "DARK CHOCOLATE BAR"
    assert values[:3] == original[:3]
    assert values[4:] == original[4:]
