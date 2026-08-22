from pathlib import Path

import pytest
from openpyxl import Workbook

from api.tools.item_combination_cleanup.workbook import (
    WorkbookError,
    WorkbookValidationError,
    export_identity_copy,
    export_approved_proposals,
    inspect_workbook,
    validate_identity_copy,
)
from tests.fixtures.item_combination_cleanup import mutate_column, write_fixture_workbook


def test_inspect_accepts_valid_workbook(tmp_path: Path):
    path = write_fixture_workbook(tmp_path / "master.xlsx")
    meta = inspect_workbook(path)
    assert meta.sheet_name == "Item_Master_List"
    assert meta.data_row_count == 5
    assert "Item_Id" in meta.headers
    assert "ExtraNote" in meta.headers
    assert meta.headers[0] == "Item_Id"


def test_inspect_allows_duplicate_barcode(tmp_path: Path):
    path = write_fixture_workbook(tmp_path / "master.xlsx")
    meta = inspect_workbook(path)
    assert meta.data_row_count == 5


def test_inspect_rejects_missing_required_column(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Item_Id", "Barcode", "Item Name", "Brand", "Size"])
    ws.append(["1", "A", "X", None, None])
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    with pytest.raises(WorkbookError, match="HsnCode"):
        inspect_workbook(path)


def test_inspect_rejects_duplicate_item_id(tmp_path: Path):
    path = write_fixture_workbook(tmp_path / "master.xlsx")
    mutate_column(path, "Item_Id", 3, "1001")
    with pytest.raises(WorkbookError, match="Duplicate Item_Id"):
        inspect_workbook(path)


def test_identity_copy_preserves_protected_cells(tmp_path: Path):
    source = write_fixture_workbook(tmp_path / "source.xlsx")
    dest = tmp_path / "output.xlsx"
    export_identity_copy(source, dest)
    validate_identity_copy(source, dest)

    from openpyxl import load_workbook
    src = load_workbook(source)
    out = load_workbook(dest)
    assert src.sheetnames == out.sheetnames
    sw, ow = src.active, out.active
    assert sw.max_row == ow.max_row
    assert [c.value for c in sw[1]] == [c.value for c in ow[1]]
    for row in range(2, sw.max_row + 1):
        assert sw.cell(row, 1).value == ow.cell(row, 1).value  # Item_Id
        assert sw.cell(row, 3).value == ow.cell(row, 3).value  # Barcode
        assert sw.cell(row, 15).value == ow.cell(row, 15).value  # MRP
        assert sw.cell(row, 16).value == ow.cell(row, 16).value  # Rate
        assert sw.cell(row, 18).value == ow.cell(row, 18).value  # ExtraNote
    assert sw.cell(6, 1).value == "001234"
    src.close()
    out.close()


def test_validator_blocks_rate_change(tmp_path: Path):
    source = write_fixture_workbook(tmp_path / "source.xlsx")
    dest = tmp_path / "output.xlsx"
    export_identity_copy(source, dest)
    mutate_column(dest, "Rate", 2, 999)
    with pytest.raises(WorkbookValidationError, match="Rate"):
        validate_identity_copy(source, dest)


def test_validator_allows_item_name_change(tmp_path: Path):
    source = write_fixture_workbook(tmp_path / "source.xlsx")
    dest = tmp_path / "output.xlsx"
    export_identity_copy(source, dest)
    mutate_column(dest, "Item Name", 2, "CHANGED NAME")
    validate_identity_copy(source, dest)


def test_export_deletes_output_when_validation_would_fail(tmp_path: Path, monkeypatch):
    source = write_fixture_workbook(tmp_path / "source.xlsx")
    dest = tmp_path / "output.xlsx"

    from api.tools.item_combination_cleanup import workbook as wbmod

    def boom(src, dst):
        raise WorkbookValidationError("protected cell changed")

    monkeypatch.setattr(wbmod, "validate_identity_copy", boom)
    with pytest.raises(WorkbookValidationError):
        wbmod.export_identity_copy(source, dest)
    assert not dest.exists()


def test_approved_export_changes_only_selected_mutable_cells(tmp_path: Path):
    source = write_fixture_workbook(tmp_path / "source.xlsx")
    dest = tmp_path / "approved.xlsx"
    export_approved_proposals(
        source,
        dest,
        {"1001": {"Item Name": "RENAMED", "Brand": "PKG", "Size": "500G"}},
    )
    from openpyxl import load_workbook
    src, out = load_workbook(source), load_workbook(dest)
    headers = {cell.value: cell.column for cell in src.active[1]}
    assert out.active.cell(2, headers["Item Name"]).value == "RENAMED"
    assert out.active.cell(2, headers["Brand"]).value == "PKG"
    assert out.active.cell(2, headers["Size"]).value == "500G"
    assert out.active.cell(2, headers["Rate"]).value == src.active.cell(2, headers["Rate"]).value
    src.close(); out.close()
